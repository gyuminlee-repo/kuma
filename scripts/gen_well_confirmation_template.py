"""Generate a synthetic, well-labeled confirmation workbook; no real assay data."""
from pathlib import Path
import shutil
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
NAME = "17_mame_well_labeled_confirmation.xlsx"


def main() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Synthetic example"
    for label, value in [("WT", 100), ("WT_2", 100), ("A1", 100), ("A1-1", 100), ("B1-1", 100)]:
        sheet.append(["Signal:", "FID1B"])
        sheet.append(["Area", "Sample Name"])
        sheet.append([value, label])
        sheet.append(["Sum", value])
        sheet.append([])
    note = book.create_sheet("Read me")
    note.append(["Synthetic parsing example only. Values are arbitrary placeholders."])
    note.append(["Well labels require matching layout or NGS verdict identities. This file does not provide them."])
    note.append(["A positive -N suffix means an additional replicate of that exact canonical label."])
    note.append(["Do not mix well labels, numeric IDs, and variant labels in one report."])
    target = ROOT / "templates" / NAME
    book.save(target)
    shutil.copyfile(target, ROOT / "src-tauri" / "samples" / "mame" / NAME)


if __name__ == "__main__":
    main()
