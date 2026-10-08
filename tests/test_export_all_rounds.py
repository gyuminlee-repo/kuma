"""A design past one plate is exported as rounds, one 96-well plate each.

The split happens in the UI, at the design-to-mappings boundary
(`src/lib/plateRounds.ts`), so `export_all` and the plate mapper only ever see
one plate per call. These tests play the part of that split with the same
chunking rule and check what the sidecar owns: the `_R1`/`_R2` folder suffix,
that each round's workbook keeps the unnumbered single-plate sheets, that no
well leaves A1-H12 or carries a `P2-` label, and that a round is refused when
its payload is not one plate or its column parity is already spent.

The boundary counts are the ones `.cross-layer-sync.json` asks for
(`plate-well-capacity`): 95, 96 and 97 straddle the split, 191, 192 and 193 the
second, 288 and 289 the third. There is no design bound any more: a round is
`R` plus its number and each round names the Echo source plate it is dispensed
from, so a design past two rounds goes onto a second 384 plate. A fixture below
96 cannot tell a split from no split.
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from pydantic import ValidationError

import sidecar_kuro.core as _core
from kuma_core.kuro.plate_mapper import _assign_well, _well_name
from sidecar_kuro.handlers.export import handle_export_all

PER_ROUND = 96
PLATE_WELLS = {_well_name(i) for i in range(PER_ROUND)}
_BASES = "ACGT"


@pytest.fixture(autouse=True)
def _reset_state():
    with _core._state_lock:
        _core._state.results = []
        _core._state.plate_mappings = []
        _core._state.dedup_info = {}
    yield
    with _core._state_lock:
        _core._state.results = []
        _core._state.plate_mappings = []
        _core._state.dedup_info = {}


def _seq(i: int, tail: str) -> str:
    # 20 nt, unique per index: the index in base 4 over a fixed stem.
    digits = ""
    n = i
    for _ in range(6):
        digits += _BASES[n % 4]
        n //= 4
    return "ATGCATGCATGC" + digits + tail


def _mutations(n: int) -> list[str]:
    return [f"M{i + 1}A" for i in range(n)]


def _round_payload(muts: list[str]) -> list[dict]:
    """One round as the UI sends it: wells re-indexed from A1 on one plate."""
    fwd = [
        {
            "well": _assign_well(i),
            "primer_name": f"{m}_F",
            "sequence": _seq(int(m[1:-1]), "AA"),
            "primer_type": "forward",
            "mutation": m,
        }
        for i, m in enumerate(muts)
    ]
    rev = [
        {
            "well": _assign_well(i),
            "primer_name": f"{m}_R",
            "sequence": _seq(int(m[1:-1]), "TT"),
            "primer_type": "reverse",
            "mutation": m,
        }
        for i, m in enumerate(muts)
    ]
    return fwd + rev


def _split(n: int) -> list[list[str]]:
    muts = _mutations(n)
    return [muts[i:i + PER_ROUND] for i in range(0, len(muts), PER_ROUND)]


def _export_round(tmp_path: Path, muts: list[str], label: str | None, *,
                  quadrant: str | None, used: list[str],
                  plate: int | None = 1) -> dict:
    params = {
        "output_dir": str(tmp_path),
        "project_name": "Proj",
        "fwd_plate_name": "F1",
        "rev_plate_name": "R1",
        "mappings": _round_payload(muts),
        "quadrant": quadrant,
        "used_quadrants": used,
    }
    if label is not None:
        params["round_label"] = label
        if plate is not None:
            params["source_plate"] = plate
    with patch("sidecar_kuro.handlers.export._dt") as m_dt:
        m_dt.now.return_value = datetime(2026, 9, 28, 10, 25)
        return handle_export_all(params)


def _all_text(folder: Path) -> str:
    """Every cell and line the round wrote, as one string to search."""
    chunks: list[str] = []
    for f in folder.iterdir():
        if f.suffix == ".xlsx":
            wb = openpyxl.load_workbook(f, read_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    chunks.extend(str(v) for v in row if v is not None)
        elif f.suffix in (".csv", ".fasta", ".json"):
            chunks.append(f.read_text(encoding="utf-8-sig"))
    return "\n".join(chunks)


def _fwd_list_wells(platemap: Path) -> list[str]:
    wb = openpyxl.load_workbook(platemap, read_only=True)
    ws = wb["Fwd List"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    idx = header.index("Well")
    return [str(r[idx]) for r in rows[1:] if r[idx] is not None]


def _plan(rounds: int) -> list[tuple[int, str, list[str]]]:
    """Plate, parity and parities already spent on that plate, per round.

    Two rounds fill one 384 source plate, so round k goes to plate k // 2 + 1.
    This is a fixture choice: the operator picks both values in the app.
    """
    plan = []
    for k in range(rounds):
        plate = k // 2 + 1
        parity = ("A1", "A2")[k % 2]
        used = [q for (p, q, _) in plan if p == plate]
        plan.append((plate, parity, used))
    return plan


@pytest.mark.parametrize(
    "n,rounds",
    [(95, 1), (96, 1), (97, 2), (191, 2), (192, 2), (193, 3), (288, 3), (289, 4)],
)
def test_split_rounds_each_land_on_one_plate(tmp_path, n, rounds):
    chunks = _split(n)
    assert len(chunks) == rounds
    assert all(len(c) <= PER_ROUND for c in chunks)
    assert sum(len(c) for c in chunks) == n

    plan = _plan(rounds)
    for k, muts in enumerate(chunks):
        label = f"R{k + 1}" if rounds > 1 else None
        plate, parity, used = plan[k]
        res = _export_round(
            tmp_path, muts, label,
            quadrant=parity if rounds > 1 else None,
            used=used if rounds > 1 else [],
            plate=plate,
        )
        assert res["failed"] == [], res["failed"]
        folder = Path(res["output_dir"])
        expected = "Proj_20260928" + (f"_{label}" if label else "")
        assert folder.name == expected

        platemap = folder / f"{folder.name}_platemap.xlsx"
        sheets = openpyxl.load_workbook(platemap, read_only=True).sheetnames
        # One plate, so no plate number on any sheet.
        assert "Fwd List" in sheets and "Fwd Plate" in sheets
        assert not any(s[-1].isdigit() for s in sheets), sheets

        wells = _fwd_list_wells(platemap)
        assert len(wells) == len(muts)
        assert set(wells) <= PLATE_WELLS
        assert "P2-" not in _all_text(folder)

        with open(folder / f"{folder.name}_echo.csv", encoding="utf-8-sig") as fh:
            names = {r["Source Plate Name"] for r in csv.DictReader(fh)}
        assert names == {f"Source [{plate if rounds > 1 else 1}]"}

    written = sorted(p.name for p in tmp_path.iterdir() if p.is_dir())
    if rounds == 1:
        assert written == ["Proj_20260928"]
    else:
        assert written == sorted(f"Proj_20260928_R{k + 1}" for k in range(rounds))


def test_a_third_round_is_accepted(tmp_path):
    # 193 needs a third chunk, which goes onto a second source plate.
    chunks = _split(193)
    assert len(chunks) == 3
    res = _export_round(tmp_path, chunks[2], "R3", quadrant="A1", used=[], plate=2)
    assert res["failed"] == [], res["failed"]
    assert Path(res["output_dir"]).name == "Proj_20260928_R3"


@pytest.mark.parametrize("label", ["R0", "R", "R01", "X1", "r3", "R-1", "R1a"])
def test_round_label_must_be_r_and_a_positive_integer(tmp_path, label):
    with pytest.raises(ValidationError):
        _export_round(tmp_path, _split(97)[0], label, quadrant="A1", used=[])
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("plate", [0, -1])
def test_source_plate_must_be_positive(tmp_path, plate):
    with pytest.raises(ValidationError):
        _export_round(tmp_path, _split(97)[0], "R1", quadrant="A1", used=[],
                      plate=plate)
    assert list(tmp_path.iterdir()) == []


def test_round_without_source_plate_is_refused(tmp_path):
    with pytest.raises(ValueError, match="choose the Echo source plate number"):
        _export_round(tmp_path, _split(97)[0], "R1", quadrant="A1", used=[],
                      plate=None)
    assert list(tmp_path.iterdir()) == []


def test_second_plate_a1_is_free_whatever_plate_one_used(tmp_path):
    # Plate 1 has A1 spent. Round 3 is on plate 2, a new plate, so the UI sends
    # nothing as used for it and A1 is allowed.
    r1, r2, r3 = _split(193)
    _export_round(tmp_path, r1, "R1", quadrant="A1", used=[], plate=1)
    res = _export_round(tmp_path, r3, "R3", quadrant="A1", used=[], plate=2)
    assert res["failed"] == []
    folder = Path(res["output_dir"])
    wb = openpyxl.load_workbook(folder / f"{folder.name}_echo.xlsx", read_only=True)
    assert wb["layout"].cell(2, 1).value == "Source [2]"


def test_same_parity_twice_on_a_later_plate_is_refused(tmp_path):
    chunks = _split(289)
    _export_round(tmp_path, chunks[2], "R3", quadrant="A2", used=[], plate=2)
    with pytest.raises(ValueError, match="already used"):
        _export_round(tmp_path, chunks[3], "R4", quadrant="A2", used=["A2"], plate=2)
    assert not (tmp_path / "Proj_20260928_R4").exists()


def test_plate_one_round_echo_matches_single_plate_bytes(tmp_path):
    # Plate 1 keeps every file byte the single-plate export writes for Echo.
    muts = _split(96)[0]
    single = Path(_export_round(tmp_path / "a", muts, None, quadrant="A1", used=[])["output_dir"])
    round1 = Path(_export_round(tmp_path / "b", muts, "R1", quadrant="A1", used=[], plate=1)["output_dir"])
    assert (single / f"{single.name}_echo.csv").read_bytes() == (
        round1 / f"{round1.name}_echo.csv"
    ).read_bytes()
    for folder in (single, round1):
        wb = openpyxl.load_workbook(folder / f"{folder.name}_echo.xlsx", read_only=True)
        assert wb["layout"].cell(2, 1).value == "Source"


def test_round_suffix_goes_on_before_the_same_day_counter(tmp_path):
    muts = _split(97)[0]
    first = _export_round(tmp_path, muts, "R1", quadrant="A1", used=[])
    again = _export_round(tmp_path, muts, "R1", quadrant="A1", used=[])
    assert Path(first["output_dir"]).name == "Proj_20260928_R1"
    assert Path(again["output_dir"]).name == "Proj_20260928_R1_2"


def test_same_parity_twice_is_refused(tmp_path):
    r1, r2 = _split(192)
    _export_round(tmp_path, r1, "R1", quadrant="A1", used=[])
    with pytest.raises(ValueError, match="already used"):
        _export_round(tmp_path, r2, "R2", quadrant="A1", used=["A1"])
    assert not (tmp_path / "Proj_20260928_R2").exists()


def test_round_without_parity_is_refused(tmp_path):
    with pytest.raises(ValueError, match="choose the Echo source plate"):
        _export_round(tmp_path, _split(97)[0], "R1", quadrant=None, used=[])
    assert list(tmp_path.iterdir()) == []


def test_unsplit_payload_is_refused_as_a_round(tmp_path):
    # 192 mappings labelled the way the overflow chunker labels them: the
    # second plate is P2-A1 .. P2-H12. That is the path a round must never take.
    muts = _mutations(192)
    params = {
        "output_dir": str(tmp_path),
        "project_name": "Proj",
        "mappings": _round_payload(muts),
        "quadrant": "A1",
        "used_quadrants": [],
        "round_label": "R1",
    }
    assert any(m["well"].startswith("P2-") for m in params["mappings"])
    with pytest.raises(ValueError, match="not on one 96-well plate"):
        handle_export_all(params)
    assert list(tmp_path.iterdir()) == []


def test_merged_rounds_reindexed_from_a1_are_refused(tmp_path):
    # Two rounds each re-indexed from A1 and then concatenated: every well is
    # on the plate, but each one is used twice.
    r1, r2 = _split(192)
    mappings = _round_payload(r1) + _round_payload(r2)
    with pytest.raises(ValueError, match="used twice"):
        handle_export_all({
            "output_dir": str(tmp_path),
            "project_name": "Proj",
            "mappings": mappings,
            "quadrant": "A1",
            "used_quadrants": [],
            "round_label": "R1",
        })
    assert list(tmp_path.iterdir()) == []


def test_single_plate_export_ignores_round_checks(tmp_path):
    # No round_label: the pre-existing export, including the no-parity layout.
    res = _export_round(tmp_path, _split(96)[0], None, quadrant=None, used=[])
    assert res["failed"] == []
    assert Path(res["output_dir"]).name == "Proj_20260928"


def test_echo_source_wells_follow_each_round_parity(tmp_path):
    r1, r2 = _split(192)
    out = {}
    for label, muts, q, used in (("R1", r1, "A1", []), ("R2", r2, "A2", ["A1"])):
        res = _export_round(tmp_path, muts, label, quadrant=q, used=used)
        folder = Path(res["output_dir"])
        with open(folder / f"{folder.name}_echo.csv", encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
        out[label] = {r["Source Well"] for r in rows}
    cols = {label: {int(w[1:]) % 2 for w in wells} for label, wells in out.items()}
    assert cols == {"R1": {1}, "R2": {0}}
    assert out["R1"].isdisjoint(out["R2"])
