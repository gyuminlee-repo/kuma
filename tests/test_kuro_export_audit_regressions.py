import csv
from contextlib import closing
from pathlib import Path

import pytest
from openpyxl import load_workbook

from kuma_core.kuro.evolvepro import load_evolvepro_csv
from kuma_core.kuro.plate_mapper import PlateMapping, export_echo_mapping_xlsx


@pytest.mark.parametrize("variant", ["M1A:A2V", "A2V:M1A", "M1A/A2V", "M1A A2V"])
def test_start_codon_combo_is_excluded(tmp_path: Path, variant: str) -> None:
    path = tmp_path / "variants.csv"
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["variant", "y_pred"])
        writer.writerows([[variant, 10], ["A2V", 1]])

    result = load_evolvepro_csv(path)

    assert result["variants"] == ["A2V"]
    assert result["start_codon_removed"] == 1


@pytest.mark.parametrize("quadrant", [None, "A1", "A13"])
@pytest.mark.parametrize("mapping_range", [None, ("C", "F")])
def test_echo_layout_matches_transfer_coordinates(
    tmp_path: Path, quadrant: str | None, mapping_range: tuple[str, str] | None,
) -> None:
    forward = PlateMapping("A1", "A2V_F", "ACGT", "forward", "A2V")
    reverse = PlateMapping("A1", "A2V_R", "TGCA", "reverse", "A2V")
    path = tmp_path / "echo.xlsx"

    export_echo_mapping_xlsx([forward], [reverse], path, quadrant=quadrant, mapping_range=mapping_range)

    with closing(load_workbook(path, read_only=True, data_only=True)) as workbook:
        layout = workbook["layout"]
        positions = {
            str(layout.cell(row=row, column=2).value) + str(layout.cell(row=4, column=col).value):
            layout.cell(row=row, column=col).value
            for row in range(5, 21) for col in range(3, 27)
            if layout.cell(row=row, column=col).value is not None
        }
        for transfer in workbook["Echo mapping file"].iter_rows(min_row=2, values_only=True):
            assert positions.get(str(transfer[2])) == transfer[1]
