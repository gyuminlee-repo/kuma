import json
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from kuma_core.mame.cli import _dump_verdicts, _load_verdicts
from kuma_core.mame.export import WellMapper, write_excel
from kuma_core.mame.models import (
    BarcodeRecord,
    NoisyPosition,
    ReplicateResult,
    TranslatedRecord,
    VerdictClass,
    VerdictRecord,
)


@pytest.mark.parametrize("scenario", ["current", "legacy", "multiple", "ambiguous"])
@pytest.mark.parametrize("measurement", [None, 0.0, 0.73])
def test_cli_roundtrip_preserves_selected_well(
    tmp_path: Path, scenario: str, measurement: float | None,
) -> None:
    # Given a selected non-PASS fallback with measured or unknown evidence.
    barcode = BarcodeRecord(
        native_barcode="NB01", custom_barcode="1_1", consensus_seq="ATG",
        file_size_kb=60.0, source_path=tmp_path / "input.fasta",
        read_count=123, n_input_reads=150, n_aligned_reads=130,
        n_mapq_failed=12, n_span_failed=8, n_mixed_positions=2,
        max_minor_allele_fraction=0.27, n_low_depth_positions=3,
        consensus_n_fraction=0.04, n_low_quality_bases=5,
        min_variant_support=measurement, n_variant_positions=2,
        min_variant_support_depth=100, median_minor_allele_fraction=0.07,
        max_minor_allele_strand_share=measurement,
        max_minor_allele_plus_count=0, max_minor_allele_minus_count=27,
        n_eligible_positions=3,
        noisy_positions=(NoisyPosition(12, 0.27, 100, 0, 27),),
        n_indel_event_positions=4, max_indel_event_fraction=0.31,
        depth_cv=measurement, depth_p10=measurement,
        depth_min_covered=None if measurement is None else int(measurement * 100),
        breadth_at_mix_min_depth=measurement, consensus_identity=measurement,
        max_del_run_length=3,
        consensus_net_indel_bp=None if measurement is None else 0,
        median_read_net_indel_bp=None if measurement is None else -1,
    )
    translated = TranslatedRecord(
        barcode=barcode, aa_sequence="M", observed_nt_changes=["T267G"],
        observed_aa_changes=["F89W"], n_no_call_aa=2, length_true_nt="ATG",
    )
    verdict = VerdictRecord(
        translated=translated, expected_mutations=["F89W"],
        verdict=VerdictClass.AMBIGUOUS, verdict_notes="indel evidence",
    )
    replicate = ReplicateResult(
        mutant_id="F89W", plate_verdicts={"NB01": verdict}, selected_plate="NB01",
        selection_reason="audit", failed=False,
        is_fallback=True, fallback_reason="no PASS replicate",
    )
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    saved = tmp_path / "verdicts.json"
    verdicts = [verdict]
    if scenario in {"multiple", "ambiguous"}:
        verdicts.append(replace(verdict, translated=replace(
            translated, barcode=replace(barcode, custom_barcode="2_1"),
        )))
    write_excel(verdicts, [replicate], first, WellMapper(), mode="amplicon")
    _dump_verdicts(verdicts, [replicate], saved)
    if scenario in {"legacy", "ambiguous"}:
        payload = json.loads(saved.read_text())
        payload["replicates"][0].pop("plate_barcodes", None)
        saved.write_text(json.dumps(payload))
    # When the actual CLI regenerates the workbook from its persisted dump.
    exported = subprocess.run(
        [sys.executable, "-m", "kuma_core.mame.cli", "export",
         "--verdict-json", str(saved), "--output", str(second)],
        capture_output=True, text=True, check=False,
    )
    if scenario == "ambiguous":
        assert exported.returncode != 0
        assert "Cannot uniquely restore" in exported.stderr
        assert not second.exists()
        return
    assert exported.returncode == 0, exported.stderr
    # Then all evidence cells, not only the selected-well row, agree.
    first_book = load_workbook(first)
    second_book = load_workbook(second)
    try:
        before = list(first_book["Final"].values)[1]
        after = list(second_book["Final"].values)[1]
        assert before[3] == "F89W"
        assert after == before, "CLI re-export must preserve the selected well and verdict"
        assert second_book.sheetnames == first_book.sheetnames
        for sheet in first_book.sheetnames:
            if sheet == "__kuma_meta__":
                continue  # The generation timestamp intentionally changes.
            assert list(second_book[sheet].values) == list(first_book[sheet].values), sheet
        restored, picks = _load_verdicts(saved)
        assert restored[0] == verdict
        assert picks[0] == replicate
    finally:
        first_book.close()
        second_book.close()


def test_cli_help_and_invalid_input(tmp_path: Path) -> None:
    command = [sys.executable, "-m", "kuma_core.mame.cli"]
    help_result = subprocess.run(command + ["--help"], capture_output=True, text=True, check=False)
    assert help_result.returncode == 0
    assert "export" in help_result.stdout
    output = tmp_path / "invalid.xlsx"
    invalid = subprocess.run(
        command + ["export", "--output", str(output)],
        capture_output=True, text=True, check=False,
    )
    assert invalid.returncode == 2
    assert not output.exists()
