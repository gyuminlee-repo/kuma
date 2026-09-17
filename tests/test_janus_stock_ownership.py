import csv
from contextlib import closing
from dataclasses import asdict, replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from kuma_core.kuro.plate_mapper import (
    PlateMapping,
    deduplicate_reverse,
    generate_plate_map,
    janus_row_values,
)
from sidecar_kuro.handlers.export import (
    handle_export_excel,
    handle_export_janus_mapping_dry_run,
    handle_export_mapping,
)
from tests.test_echo_multiplate_ownership import ownership_results


def stock_case(case: str) -> tuple[list[PlateMapping], dict[str, list[str]]]:
    if case == "sparse":
        return [
            PlateMapping("B2", "M1_F", "AAAA", "forward", "M1"),
            PlateMapping("D1", "M2_F", "CCCC", "forward", "M2"),
            PlateMapping("A1", "M1_R", "TTTT", "reverse", "M1"),
            PlateMapping("B1", "M2_R", "GGGG", "reverse", "M2"),
        ], {"TTTT": ["M1"], "GGGG": ["M2"]}
    count = {"default": 5, "multiplate": 193, "reduced": 97, "row": 49}[case]
    results = ownership_results(count)
    if case == "default":
        results = [replace(r, reverse_seq="ACGT" + "A" * i) for i, r in enumerate(results)]
    fwd, rev = generate_plate_map(
        results, well_order="row" if case == "row" else "column",
        mapping_range=("C", "J") if case in ("reduced", "row") else None,
    )
    return fwd + rev, deduplicate_reverse(results)


@pytest.mark.parametrize("case", ["sparse", "default", "multiplate", "reduced", "row"])
def test_public_janus_exports_aspirate_ordered_stock(tmp_path: Path, case: str) -> None:
    mappings, groups = stock_case(case)
    payload = {"mappings": [asdict(m) for m in mappings], "dedup_info": groups}
    expected_reverse = {mutation: sequence for sequence, mutations in groups.items() for mutation in mutations}
    expected_forward = {m.mutation: m.sequence for m in mappings if m.primer_type == "forward"}
    order_path = tmp_path / "order.xlsx"

    assert handle_export_excel({**payload, "filepath": str(order_path)})["success"]
    preview = handle_export_janus_mapping_dry_run(payload)
    rows = preview["rows"]
    for extension in ("csv", "xlsx"):
        assert handle_export_mapping({
            **payload, "format": "janus", "filepath": str(tmp_path / f"janus.{extension}"),
        })["success"]

    stock: dict[tuple[str, str], tuple[str, str]] = {}
    with closing(load_workbook(order_path)) as workbook:
        for sheet in workbook:
            if " List" not in sheet.title:
                continue
            plate = int(sheet.title.split()[-1]) if sheet.title[-1].isdigit() else 1
            role = "fwd" if sheet.title.startswith("Fwd") else "rev"
            for well, name, sequence, *_ in list(sheet.values)[1:]:
                address = str(well) if plate == 1 else f"P{plate}-{well}"
                stock[(role, address)] = (str(name), str(sequence))
    assert len(rows) == preview["total"] == len(expected_forward) * 2
    for row in rows:
        expected = expected_forward if row["role"] == "fwd" else expected_reverse
        assert stock.get((row["role"], row["asp_posi"]), ("", ""))[1] == expected[row["mutation"]], row
        assert row["dsp_posi"] == next(m.well for m in mappings if m.primer_type == "forward" and m.mutation == row["mutation"])
    if case == "sparse":
        assert [r["asp_posi"] for r in rows if r["role"] == "rev"] == ["D1", "B2"]
    if case == "default":
        assert rows[-1]["asp_posi"] == "E1"

    with (tmp_path / "janus.csv").open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.reader(stream))[1:]
    for actual, row in zip(csv_rows, rows, strict=True):
        assert actual[:7] == [str(v) for v in janus_row_values(row)[:7]]
        assert float(actual[7]) == row["volume"]
    with closing(load_workbook(tmp_path / "janus.xlsx")) as workbook:
        assert list(workbook["primer_mapping file"].values)[1:] == [tuple(janus_row_values(r)) for r in rows]
        for (role, address), (name, _) in stock.items():
            prefix, separator, base = address.partition("-")
            plate = int(prefix[1:]) if separator else 1
            well = base if separator else address
            sheet = workbook["layout" if plate == 1 else f"layout {plate}"]
            start = 4 if role == "fwd" else 15
            assert sheet.cell(start + ord(well[0]) - ord("A"), 2 + int(well[1:])).value == name
            if role == "rev":
                usage = [r for r in sheet.iter_rows(min_row=37, values_only=True) if r[0] == well and r[1] == name]
                assert len(usage) == 1
                assert usage[0][2] == sum(r["role"] == "rev" and r["asp_posi"] == address for r in rows)
