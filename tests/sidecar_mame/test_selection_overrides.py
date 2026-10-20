"""Manual plate-copy confirmation: override application and the files it writes."""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pytest

from kuma_core.mame.models import (
    BarcodeRecord,
    ReplicateResult,
    TranslatedRecord,
    VerdictClass,
    VerdictRecord,
)
from kuma_core.mame.select.manual_override import apply_manual_overrides
from sidecar_mame.core import get_state, reset_state, set_last_analyze
from sidecar_mame.handlers.analyze import (
    _deserialize_replicate,
    _serialize_replicate,
    _serialize_verdict,
    picks_autosave_path,
)
from sidecar_mame.handlers.load import handle_load_analyze_result
from sidecar_mame.handlers.export import (
    handle_apply_selection_overrides,
    handle_export_janus_mapping,
    handle_export_janus_mapping_dry_run,
    handle_get_plate_data,
)


def _verdict(nb: str, custom: str, verdict: VerdictClass) -> VerdictRecord:
    barcode = BarcodeRecord(
        native_barcode=nb,
        custom_barcode=custom,
        consensus_seq="",
        file_size_kb=10.0,
        source_path=Path("/tmp/mock.fasta"),
    )
    translated = TranslatedRecord(
        barcode=barcode, aa_sequence="", observed_aa_changes=[], observed_nt_changes=[]
    )
    return VerdictRecord(
        translated=translated,
        expected_mutations=[],
        verdict=verdict,
        verdict_notes="",
        mutant_id="",
    )


def _replicate(mutant: str, custom: str, copies: dict[str, VerdictClass], selected: str):
    verdicts = {nb: _verdict(nb, custom, cls) for nb, cls in copies.items()}
    return ReplicateResult(
        mutant_id=mutant,
        plate_verdicts=verdicts,
        selected_plate=selected,
        selection_reason=f"auto {selected}",
    )


@pytest.fixture
def run(tmp_path: Path):
    """B1 holds a shallow selected copy and a PASS copy; A1 is a plain PASS."""
    a1 = _replicate("A1MUT", "1_1", {"NB07": VerdictClass.PASS}, "NB07")
    b1 = _replicate(
        "B1MUT",
        "2_1",
        {"NB07": VerdictClass.LOWDEPTH, "NB08": VerdictClass.PASS},
        "NB07",
    )
    verdicts = [vr for rr in (a1, b1) for vr in rr.plate_verdicts.values()]
    output = tmp_path / "run.xlsx"
    set_last_analyze(verdicts, [a1, b1], str(output), run_meta=None)
    yield output
    reset_state()


def _cached() -> list[ReplicateResult]:
    return get_state().last_replicates or []


def _override(plate: str = "NB08", note: str = "Sanger OK") -> dict:
    return {"overrides": [{"mutant_id": "B1MUT", "plate": plate, "note": note}]}


def _picks(output: Path) -> list[dict[str, str]]:
    with picks_autosave_path(output).open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def test_override_swaps_copy_and_keeps_the_algorithm_verdict(run) -> None:
    reps = _cached()
    out, applied, dropped = apply_manual_overrides(
        reps, _override()["overrides"]
    )
    b1 = out[1]
    assert dropped == []
    assert (b1.selected_plate, b1.failed, b1.is_fallback) == ("NB08", False, False)
    assert b1.manual is not None
    assert (b1.manual.auto_selected_plate, b1.manual.auto_verdict) == ("NB07", "LOWDEPTH")
    assert applied[0]["auto_verdict"] == "LOWDEPTH"
    assert reps[1].selected_plate == "NB07"
    assert out[0].manual is None


def test_empty_list_restores_the_automatic_pick(run) -> None:
    handle_apply_selection_overrides(_override())
    handle_apply_selection_overrides({"overrides": []})
    b1 = _cached()[1]
    assert (b1.selected_plate, b1.selection_reason, b1.manual) == ("NB07", "auto NB07", None)


def test_unusable_entries_are_dropped_with_a_reason(run) -> None:
    result = handle_apply_selection_overrides(
        {
            "overrides": [
                {"mutant_id": "NOPE", "plate": "NB08", "note": ""},
                {"mutant_id": "B1MUT", "plate": "NB99", "note": ""},
                {"mutant_id": "B1MUT", "plate": "NB08", "note": "ok"},
                {"mutant_id": "B1MUT", "plate": "NB07", "note": "again"},
            ]
        }
    )
    assert [d["reason"] for d in result["dropped"]] == [
        "unknown_mutant",
        "unknown_copy",
        "duplicate_mutant",
    ]
    assert [a["plate"] for a in result["applied"]] == ["NB08"]


def test_pick_list_follows_the_override_and_records_it(run) -> None:
    handle_apply_selection_overrides(_override(plate="NB07", note="Sanger 2026-10-20 OK"))
    rows = {r["name"]: r for r in _picks(run)}
    assert set(rows) == {"A1MUT", "B1MUT"}
    assert rows["B1MUT"]["manual_confirmation"] == "Sanger 2026-10-20 OK"
    assert rows["A1MUT"]["manual_confirmation"] == ""

    handle_apply_selection_overrides(_override(plate="NB08"))
    assert {r["name"]: r["source_plate"] for r in _picks(run)}["B1MUT"] == "NB08"


def test_pick_list_columns_are_unchanged_without_an_override(run) -> None:
    result = handle_apply_selection_overrides({"overrides": []})
    assert result["picks_autosave"]["status"] == "saved"
    rows = _picks(run)
    assert list(rows[0]) == [
        "name",
        "source_plate",
        "source_well",
        "dest_well",
        "priority_score",
    ]
    assert [r["name"] for r in rows] == ["A1MUT"]


def test_undoing_the_only_pick_removes_the_stale_pick_list(tmp_path: Path) -> None:
    only = _replicate("ONLY", "1_1", {"NB07": VerdictClass.LOWDEPTH}, "NB07")
    output = tmp_path / "solo.xlsx"
    set_last_analyze(list(only.plate_verdicts.values()), [only], str(output), run_meta=None)
    try:
        handle_apply_selection_overrides(
            {"overrides": [{"mutant_id": "ONLY", "plate": "NB07", "note": "ok"}]}
        )
        assert picks_autosave_path(output).exists()
        handle_apply_selection_overrides({"overrides": []})
        assert not picks_autosave_path(output).exists()
    finally:
        reset_state()


def test_plate_data_marks_the_confirmed_copy(run) -> None:
    handle_apply_selection_overrides(_override(note="repeat run"))
    wells = {(w["native_barcode"], w["well"]): w for w in handle_get_plate_data({})["wells"]}
    chosen = wells[("NB08", "B1")]
    assert (chosen["selected"], chosen["manually_confirmed"], chosen["manual_note"]) == (
        True,
        True,
        "repeat run",
    )
    assert wells[("NB07", "B1")]["selected"] is False
    assert wells[("NB07", "A1")]["manually_confirmed"] is False


def test_janus_export_carries_the_confirmation(run, tmp_path: Path) -> None:
    handle_apply_selection_overrides(_override(plate="NB07", note="Sanger OK"))
    preview = handle_export_janus_mapping_dry_run({"output_schema": "legacy5"})
    assert [w["code"] for w in preview["warnings"]] == ["manual_confirmation"]

    out = tmp_path / "janus.xlsx"
    handle_export_janus_mapping(
        {"output": str(out), "format": "xlsx", "output_schema": "legacy5"}
    )
    wb = openpyxl.load_workbook(out)
    sheet = wb["Janus Mapping"]
    assert [c.value for c in sheet[1]][-1] == "manual_confirmation"
    assert [c.value for c in sheet[3]][-1] == "Sanger OK"
    assert ("manual_confirmation", "B1MUT: Sanger OK") in [
        tuple(c.value for c in row) for row in wb["__kuma_meta__"].iter_rows()
    ]

    device = tmp_path / "device.csv"
    handle_export_janus_mapping({"output": str(device)})
    with device.open(encoding="utf-8") as fh:
        assert len(next(csv.reader(fh))) == 8


def test_serialized_replicate_round_trips_and_omits_the_key_when_automatic(run) -> None:
    handle_apply_selection_overrides(_override())
    first, confirmed = _cached()
    assert "manual" not in _serialize_replicate(first)
    rebuilt = _deserialize_replicate(_serialize_replicate(confirmed))
    assert rebuilt.manual == confirmed.manual
    assert rebuilt.selected_plate == "NB08"


def _final_rows(workbook: Path) -> tuple[list[str], dict[str, list[object]]]:
    sheet = openpyxl.load_workbook(workbook)["Final"]
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    return [str(v) for v in rows[0]], {str(r[0]): r for r in rows[1:] if r[0]}


def test_result_workbook_on_disk_follows_the_override_and_the_undo(run) -> None:
    result = handle_apply_selection_overrides(_override(note="Sanger OK"))
    assert result["workbook_autosave"]["status"] == "saved"
    header, rows = _final_rows(run)
    assert header[-1] == "manual_confirmation"
    assert rows["B1"][1] == "NB08"
    assert rows["B1"][-1] == "Sanger OK"

    handle_apply_selection_overrides({"overrides": []})
    header, rows = _final_rows(run)
    assert "manual_confirmation" not in header
    assert rows["B1"][1] == "NB07"


def test_unwritable_workbook_is_reported_and_keeps_the_confirmation(
    run, tmp_path: Path
) -> None:
    state = get_state()
    set_last_analyze(
        state.last_verdicts or [],
        state.last_replicates or [],
        str(tmp_path / "missing_dir" / "run.xlsx"),
        run_meta=None,
    )
    result = handle_apply_selection_overrides(_override())
    assert result["workbook_autosave"]["status"] == "failed"
    assert _cached()[1].selected_plate == "NB08"


def test_confirmation_survives_the_result_file_round_trip(run) -> None:
    handle_apply_selection_overrides(_override(note="Sanger OK"))
    state = get_state()
    payload = {
        "verdicts": [_serialize_verdict(v) for v in state.last_verdicts or []],
        "replicates": [_serialize_replicate(r) for r in _cached()],
        "output_path": str(run),
    }

    reset_state()
    handle_load_analyze_result(payload)
    confirmed = _cached()[1]
    assert (confirmed.selected_plate, confirmed.manual.note if confirmed.manual else None) == (
        "NB08",
        "Sanger OK",
    )

    handle_apply_selection_overrides({"overrides": []})
    restored = _cached()[1]
    assert (restored.selected_plate, restored.manual) == ("NB07", None)
