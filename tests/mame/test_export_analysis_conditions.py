"""A MAME workbook has to state the conditions it was produced under.

A finished result names plates, wells and verdicts, and says nothing about the
reference it graded them against or the thresholds that turned reads into a
verdict. Two workbooks produced from the same plate under a different
``min_read_count`` or a different reference file are indistinguishable once the
session that made them is gone, and the person who has to decide whether a
result can be compared with last month's has no way to answer it from the file.

These cases read the workbook back the way that person would and require the
conditions to be in it, as the values the run was actually called with.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pytest

from kuma_core.mame.models import BarcodeRecord, ExpectedMutation

# A 12 nt CDS (M K V F) so a whole run fits in the fixture.
_REFERENCE_SEQ = "ATGAAAGTTTTT"
_REFERENCE_SHA256 = hashlib.sha256(_REFERENCE_SEQ.encode("ascii")).hexdigest()

_FINAL_SUMMARY_CONTENT = """\
instrument = PAX12345
position = X3
flow_cell_id = FAW12345
sample_id = my_sample
kit = SQK-LSK109
started = 2024-03-15T10:00:00Z
basecalling_enabled = true
"""


def _record(custom_barcode: str, seq: str) -> BarcodeRecord:
    return BarcodeRecord(
        native_barcode="NB01",
        custom_barcode=custom_barcode,
        consensus_seq=seq,
        file_size_kb=10.0,
        source_path=Path(f"/tmp/{custom_barcode}.fasta"),
        read_count=100,
    )


def _expected() -> list[ExpectedMutation]:
    return [
        ExpectedMutation(
            mutant_id="K2R", position=2, wt_aa="K", mt_aa="R",
            wt_codon="AAA", mt_codon="CGT", group_id="", primer_set_ref="K2R",
            notation_type="substitution", status="DESIGNED",
        ),
    ]


def _meta_rows(xlsx: Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(xlsx)
    try:
        ws = wb["__kuma_meta__"]
        return {
            str(row[0]): ("" if row[1] is None else str(row[1]))
            for row in ws.iter_rows(min_row=2, max_col=2, values_only=True)
            if row[0]
        }
    finally:
        wb.close()


def _run(tmp_path: Path, input_dir: Path | None = None, **overrides: object) -> Path:
    from kuma_core.mame.pipeline import run_analyze

    reference = tmp_path / "ispS_amplicon.fasta"
    reference.write_text(f">ref\n{_REFERENCE_SEQ}\n", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    kwargs: dict[str, object] = dict(
        input_dir=input_dir if input_dir is not None else tmp_path,
        reference_path=reference,
        expected_path=tmp_path / "unused.xlsx",
        output_path=out,
        cds_start=0,
        cds_end=12,
        mode="amplicon",
        min_file_size_kb=50.0,
        min_read_count=30,
        max_consensus_n_fraction=0.0,
        many_cutoff=5,
        records=[_record("1_1", "ATGCGTGTTTTT")],
        expected_mutations=_expected(),
    )
    kwargs.update(overrides)
    run_analyze(**kwargs)  # type: ignore[arg-type]
    return out


def test_the_workbook_states_the_conditions_it_was_produced_under(
    tmp_path: Path,
) -> None:
    """Every parameter the run was called with is readable from the file.

    Literals rather than a second call into the writer: comparing the sheet
    against the same code that filled it would assert nothing.
    """
    pytest.importorskip("openpyxl")

    kv = _meta_rows(_run(tmp_path))

    assert kv.get("reference_file") == "ispS_amplicon.fasta"
    assert kv.get("reference_length") == "12"
    assert kv.get("reference_sha256") == _REFERENCE_SHA256
    assert kv.get("coding_window") == "0-12"
    assert kv.get("mode") == "amplicon"
    assert kv.get("ingest_mode") == "barcode"
    assert kv.get("min_read_count") == "30"
    assert kv.get("max_consensus_n_fraction") == "0.0"
    assert kv.get("min_file_size_kb") == "50.0"
    assert kv.get("many_cutoff") == "5"
    # 30 x 3: below this depth a minor allele is not told apart from ONT error.
    assert kv.get("mixed_confident_depth_factor") == "3"
    assert kv.get("mixed_confident_read_count") == "90"


def test_the_workbook_names_the_vocabulary_its_verdicts_were_drawn_from(
    tmp_path: Path,
) -> None:
    """The reader must not have to guess which classes were in play.

    Derived from ``VerdictClass`` rather than spelled out, because a
    hand-written expectation here is the same hand-copy that drifted from the
    definition downstream and hid whole wells. The enum is the definition, not
    the writer, so reading it is not comparing the sheet against itself.
    """
    pytest.importorskip("openpyxl")
    from kuma_core.mame.models import VerdictClass

    kv = _meta_rows(_run(tmp_path))

    assert kv["verdict_classes"].split(", ") == [v.value for v in VerdictClass]


def test_no_condition_row_is_left_blank(tmp_path: Path) -> None:
    """A key with an empty value records nothing while looking like a record."""
    pytest.importorskip("openpyxl")

    kv = _meta_rows(_run(tmp_path))
    blank = [k for k, v in kv.items() if v.strip() == ""]
    assert blank == [], f"blank __kuma_meta__ values: {blank}"


def test_thresholds_that_are_off_say_so_rather_than_going_blank(
    tmp_path: Path,
) -> None:
    """``None`` is a real setting (the filter is off), not a missing value."""
    pytest.importorskip("openpyxl")

    kv = _meta_rows(
        _run(tmp_path, min_read_count=None, max_consensus_n_fraction=None)
    )

    assert kv.get("min_read_count") == "disabled"
    assert kv.get("max_consensus_n_fraction") == "disabled"
    assert kv.get("mixed_confident_read_count") == "disabled"


def test_the_reference_hash_follows_the_sequence_not_the_file(
    tmp_path: Path,
) -> None:
    """The same molecule under a different name and line width is the same run.

    Hashing file bytes would report two runs as incomparable because someone
    re-wrapped the FASTA.
    """
    pytest.importorskip("openpyxl")

    a = tmp_path / "a"
    b = tmp_path / "b"
    a.mkdir()
    b.mkdir()
    kv_a = _meta_rows(_run(a))
    # Same sequence, different header and wrapped at 4 nt.
    other = b / "renamed.fa"
    other.write_text(">something else\nATGA\nAAGT\nTTTT\n", encoding="utf-8")
    kv_b = _meta_rows(_run(b, reference_path=other))

    assert kv_a["reference_sha256"] == kv_b["reference_sha256"]
    assert kv_b["reference_file"] == "renamed.fa"


def test_the_library_path_stamps_the_release_version(tmp_path: Path) -> None:
    """``run_analyze`` wrote workbooks with an empty ``kuma_version`` cell."""
    pytest.importorskip("openpyxl")
    from kuma_core.shared.version import KUMA_VERSION

    kv = _meta_rows(_run(tmp_path))
    assert kv.get("kuma_version") == KUMA_VERSION


def test_run_analyze_finds_the_minknow_run_it_was_pointed_at(
    tmp_path: Path,
) -> None:
    """``ngs_run_meta`` read "(not found)" on every library-path workbook.

    ``discover_run_meta`` existed and nothing on this path called it, so the
    sequencing run behind a result was recorded as absent even when the reads
    sat inside the MinKNOW folder that names it.
    """
    pytest.importorskip("openpyxl")

    run_dir = tmp_path / "run_xyz"
    run_dir.mkdir()
    (run_dir / "final_summary_PAX12345_abc.txt").write_text(
        _FINAL_SUMMARY_CONTENT, encoding="utf-8"
    )
    input_dir = run_dir / "sort_barcode06"
    input_dir.mkdir()

    kv = _meta_rows(_run(tmp_path, input_dir=input_dir))

    assert kv.get("flow_cell_id") == "FAW12345"
    assert kv.get("instrument") == "PAX12345"
    assert "ngs_run_meta" not in kv


def test_a_caller_that_already_knows_the_run_meta_is_not_made_to_rediscover(
    tmp_path: Path,
) -> None:
    """The sidecar discovers run meta on its own thread and hands it in.

    Passing it explicitly must be honoured, including an explicit ``None``,
    or the pipeline pays for a second filesystem probe of the share.
    """
    pytest.importorskip("openpyxl")
    from kuma_core.mame.ingest.run_meta import NgsRunMeta

    run_dir = tmp_path / "run_xyz"
    run_dir.mkdir()
    (run_dir / "final_summary_PAX12345_abc.txt").write_text(
        _FINAL_SUMMARY_CONTENT, encoding="utf-8"
    )
    input_dir = run_dir / "sort_barcode06"
    input_dir.mkdir()

    supplied = NgsRunMeta(
        instrument="",
        position="",
        flow_cell_id="HANDED-IN",
        sample_id="",
        kit="",
        started="",
        basecalling_enabled=None,
        raw_run_dir=None,
    )
    kv = _meta_rows(_run(tmp_path, input_dir=input_dir, ngs_run_meta=supplied))
    assert kv.get("flow_cell_id") == "HANDED-IN"

    kv_none = _meta_rows(_run(tmp_path, input_dir=input_dir, ngs_run_meta=None))
    assert "ngs_run_meta" in kv_none


def test_an_export_without_conditions_says_they_are_missing(tmp_path: Path) -> None:
    """The CLI export path has no conditions to state, and has to admit it."""
    pytest.importorskip("openpyxl")
    from kuma_core.mame.export import write_excel

    from tests.mame.test_run_meta import _make_replicate, _make_verdict

    out = tmp_path / "bare.xlsx"
    write_excel(
        verdict_records=[_make_verdict("NB01", "1_1")],
        replicate_results=[_make_replicate("V5F", "NB01", "1_1")],
        output_path=out,
    )
    kv = _meta_rows(out)
    assert "analysis_conditions" in kv
    assert kv["analysis_conditions"].startswith("(not recorded")


def _meta_keys(xlsx: Path) -> list[str]:
    """Column-A keys in sheet order, duplicates kept.

    ``_meta_rows`` builds a dict, and a dict collapses a key that was written
    twice into one entry, so it cannot answer "does this row appear exactly
    once". Counting has to read the rows as they were appended.
    """
    wb = openpyxl.load_workbook(xlsx)
    try:
        ws = wb["__kuma_meta__"]
        return [
            str(row[0])
            for row in ws.iter_rows(min_row=2, max_col=2, values_only=True)
            if row[0]
        ]
    finally:
        wb.close()


def _bare_workbook(out: Path) -> Path:
    """A workbook written by a caller that ran no analysis."""
    from kuma_core.mame.export import write_excel

    from tests.mame.test_run_meta import _make_replicate, _make_verdict

    write_excel(
        verdict_records=[_make_verdict("NB01", "1_1")],
        replicate_results=[_make_replicate("V5F", "NB01", "1_1")],
        output_path=out,
    )
    return out


def test_the_verdict_vocabulary_is_recorded_even_without_analysis_conditions(
    tmp_path: Path,
) -> None:
    """The class set is a property of kuma, not of one run.

    ``reference_sha256`` or ``min_read_count`` are true of one execution and of
    nothing else, so a workbook written by a caller that ran no analysis cannot
    state them. The eight verdict classes are the same whoever wrote the file,
    and a reader that has to hand-copy them because this workbook happens to
    come from the export path is the drift this row exists to stop.
    """
    pytest.importorskip("openpyxl")
    from kuma_core.mame.models import VerdictClass

    kv = _meta_rows(_bare_workbook(tmp_path / "bare.xlsx"))

    assert kv["verdict_classes"].split(", ") == [v.value for v in VerdictClass]
    # The other fact the sheet states on this path is untouched and still true.
    assert kv["analysis_conditions"].startswith("(not recorded")


def test_the_verdict_vocabulary_row_is_written_once_on_either_path(
    tmp_path: Path,
) -> None:
    """Two copies of the row would be two answers to one question.

    A reader that finds the key twice has to decide which cell to trust, and a
    later edit to one of the two producers would make that choice matter.
    """
    pytest.importorskip("openpyxl")

    with_analysis = _meta_keys(_run(tmp_path))
    without_analysis = _meta_keys(_bare_workbook(tmp_path / "bare.xlsx"))

    assert with_analysis.count("verdict_classes") == 1
    assert without_analysis.count("verdict_classes") == 1
