"""Membership policy for GC report rows, one pair of fixtures per slot.

Every slot reads rows of one unit (well, numeric ID, variant). A row is
(1) used when it is that unit and belongs to the run, (2) set aside with its
name and position when nothing ties it to the run, or (3) refuses the file
when it belongs to the run but is written in another notation. A file with no
row of kind (1) is refused too. Each test below pairs a mixture that carries
membership evidence (refused) with the same shape that carries none (passes),
so the line between the two is what is pinned.

All fixtures are synthetic.
"""
from __future__ import annotations

import json
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from kuma_core.mame.activity.build_evolvepro_input import OperatorExclusion, build_evolvepro_input
from kuma_core.mame.activity.evolvepro_xlsx import read_evolvepro_rows

# A01 5F, B01 10L, C01 11I, H12 WT.
PLATE = [("V5F", "A1"), ("K10L", "B1"), ("S11I", "C1"), ("WT", "H12")]


def _layout(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Mutant", "Well Pos."])
    for row in PLATE:
        ws.append(row)
    wb.save(path)
    return path


def _verdict(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["well_id", "mutant_id", "verdict"])
    for mutant, well in PLATE:
        ws.append([f"{well[0]}{int(well[1:]):02d}", mutant, "PASS"])
    wb.save(path)
    return path


def _blocks(path: Path, samples: Sequence[tuple[str, object]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    for name, area in samples:
        ws.append(["Signal:", "FID1B"])
        ws.append(["Area", "Sample Name"])
        ws.append([area, name])
        ws.append(["Sum", area])
        ws.append([])
    wb.save(path)
    return path


def _flat(path: Path, samples: Sequence[tuple[str, object]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Sample Name", "Area"])
    for row in samples:
        ws.append(list(row))
    wb.save(path)
    return path


def _cell(path: Path, row: int, column: int) -> object:
    """A cell as a spreadsheet opens it, 1-based."""
    sheet = openpyxl.load_workbook(path).active
    assert sheet is not None
    return sheet.cell(row=row, column=column).value


@pytest.fixture
def plate(tmp_path: Path) -> dict[str, Path]:
    return {"layout": _layout(tmp_path / "layout.xlsx"), "verdict": _verdict(tmp_path / "verdict.xlsx")}


def _build(tmp_path: Path, plate: dict[str, Path], **sources: Any):
    return build_evolvepro_input(
        tmp_path / "out.xlsx", layout_xlsx=plate["layout"], verdict_xlsx=plate["verdict"], **sources
    )


def _manifest(result) -> dict:
    assert result.manifest_path is not None
    return json.loads(result.manifest_path.read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# The reported symptom: a shared run whose block report carries a BLANK
# ---------------------------------------------------------------------------

SHARED_RUN = [
    ("WT1", 10.0), ("WT2", 10.0), ("WT3", 10.0),
    ("A1", 20.0), ("BLANK", 0.3), ("B1", 5.0), ("C1", 12.0),
]


def test_a_blank_in_a_well_report_is_set_aside_with_its_position(tmp_path: Path, plate) -> None:
    """Before this change the build refused: "sample 'BLANK' is not a well"."""
    result = _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", SHARED_RUN))
    assert dict(read_evolvepro_rows(tmp_path / "out.xlsx")) == pytest.approx({"5F": 2.0, "10L": 0.5, "11I": 1.2})
    assert "round1_report_xlsx excluded 1 row(s) (non_well_sample): 'BLANK' (block 5, row 23)" in result.warnings
    assert result.exclusion_reason_counts == {"non_well_sample": 1}
    manifest = _manifest(result)
    assert manifest["exclusion_reason_counts"] == {"non_well_sample": 1}
    assert manifest["excluded_rows"] == [{
        "source": "round1_report_xlsx", "sample_name": "BLANK", "reason": "non_well_sample",
        "row_index": 23, "block_index": 5, "area": 0.3, "kept": False,
    }]
    # The WT rows the reference came from, by name and position.
    assert [(row["sample_name"], row["block_index"], row["row_index"]) for row in result.wt_rows] == [
        ("WT1", 1, 3), ("WT2", 2, 8), ("WT3", 3, 13),
    ]


def test_a_blank_in_a_flat_sheet_is_set_aside_too(tmp_path: Path, plate) -> None:
    """The flat export reached the same refusal through another reader."""
    sheet = _flat(tmp_path / "gc.xlsx", [("A1", 2.0), ("BLANK", 0.01), ("5", 0.5), ("B1", 0.5)])
    result = _build(tmp_path, plate, gc_data_xlsx=sheet)
    assert dict(read_evolvepro_rows(tmp_path / "out.xlsx")) == pytest.approx({"5F": 2.0, "10L": 0.5})
    assert result.exclusion_reason_counts == {"non_well_sample": 1, "numeric_calibration": 1}
    assert {(row["sample_name"], row["row_index"], row["block_index"]) for row in _manifest(result)["excluded_rows"]} == {
        ("BLANK", 3, None), ("5", 4, None),
    }


def test_calibration_numbers_in_a_well_report_do_not_refuse(tmp_path: Path, plate) -> None:
    """Concentration-named standards (1, 5, 10) are in range of any order, and still not IDs here."""
    report = _blocks(tmp_path / "r.xlsx", [("1", 3.0), ("5", 7.0), ("10", 9.0)] + SHARED_RUN[:4])
    result = _build(tmp_path, plate, round1_report_xlsx=report)
    assert result.exclusion_reason_counts == {"numeric_calibration": 3}
    assert sum("(numeric_calibration): '1' (block 1, row 3), '5' (block 2, row 8)" in w for w in result.warnings) == 1


# ---------------------------------------------------------------------------
# Well slot (one replicate)
# ---------------------------------------------------------------------------

def test_well_slot_refuses_this_runs_variant_label(tmp_path: Path, plate) -> None:
    report = _blocks(tmp_path / "r.xlsx", SHARED_RUN[:4] + [("11I", 9.0)])
    with pytest.raises(ValueError, match=r"block 5, row 23: sample '11I' names variant 11I of this run"):
        _build(tmp_path, plate, round1_report_xlsx=report)


def test_well_slot_sets_aside_a_variant_label_the_run_did_not_make(tmp_path: Path, plate) -> None:
    report = _blocks(tmp_path / "r.xlsx", SHARED_RUN[:4] + [("99Y", 9.0)])
    result = _build(tmp_path, plate, round1_report_xlsx=report)
    assert result.exclusion_reason_counts == {"non_well_sample": 1}


def test_a_file_with_no_row_of_this_run_is_refused_on_membership(tmp_path: Path, plate) -> None:
    report = _blocks(tmp_path / "r.xlsx", [("WT1", 10.0), ("BLANK", 1.0), ("G7", 2.0)])
    with pytest.raises(ValueError, match="no row that belongs to this run.*'BLANK' \\(block 2, row 8, non_well_sample\\)"):
        _build(tmp_path, plate, round1_report_xlsx=report)


def test_rows_without_a_name_or_an_area_are_reported(tmp_path: Path, plate) -> None:
    report = _blocks(tmp_path / "r.xlsx", SHARED_RUN[:4] + [("", 4.0), ("B1", None)])
    result = _build(tmp_path, plate, round1_report_xlsx=report)
    assert result.exclusion_reason_counts == {"no_name": 1, "no_area": 1}


def test_repeats_are_flagged_and_kept(tmp_path: Path, plate) -> None:
    """A repeated WT number or a well measured twice may be another experiment's; it is said loudly."""
    report = _blocks(tmp_path / "r.xlsx", SHARED_RUN[:4] + [("WT1", 10.0), ("A1", 40.0)])
    result = _build(tmp_path, plate, round1_report_xlsx=report)
    assert result.exclusion_reason_counts == {"duplicate_replicate": 1, "duplicate_well": 1}
    assert result.variant_replicates["5F"] == pytest.approx([2.0, 4.0])
    assert any("duplicate_replicate: 1 row(s) kept in the mean but flagged" in w and "'WT1' (block 5, row 23)" in w for w in result.warnings)
    assert all(row["kept"] for row in result.excluded_rows)


# ---------------------------------------------------------------------------
# Numeric-ID slot
# ---------------------------------------------------------------------------

def test_numeric_slot_refuses_a_well_of_the_plate(tmp_path: Path, plate) -> None:
    report = _blocks(tmp_path / "n.xlsx", [("WT1", 10.0), ("1", 20.0), ("2", 5.0), ("3", 12.0), ("A1", 9.0)])
    with pytest.raises(ValueError, match="sample 'A1' names well A01 of the plate layout"):
        _build(tmp_path, plate, numeric_report_xlsx=report)


def test_numeric_slot_sets_aside_names_and_wells_of_no_membership(tmp_path: Path, plate) -> None:
    report = _blocks(
        tmp_path / "n.xlsx",
        [("WT1", 10.0), ("1", 20.0), ("BLANK", 0.1), ("2", 5.0), ("3", 12.0), ("G7", 3.0)],
    )
    result = _build(tmp_path, plate, numeric_report_xlsx=report)
    assert dict(read_evolvepro_rows(tmp_path / "out.xlsx")) == pytest.approx({"5F": 2.0, "10L": 0.5, "11I": 1.2})
    assert result.exclusion_reason_counts == {"unrecognized_sample": 2}
    assert {(row["sample_name"], row["block_index"]) for row in result.excluded_rows} == {("BLANK", 3), ("G7", 6)}


def test_numeric_slot_refuses_an_id_past_the_plate(tmp_path: Path, plate) -> None:
    """An ID is a position: one past the plate says the file or the order is wrong."""
    report = _blocks(tmp_path / "n.xlsx", [("WT1", 10.0), ("1", 20.0), ("2", 5.0), ("3", 12.0), ("200", 9.0)])
    with pytest.raises(ValueError, match=r"IDs must be 1\.\.3"):
        _build(tmp_path, plate, numeric_report_xlsx=report)


# ---------------------------------------------------------------------------
# Variant/well-labelled replicated confirmation
# ---------------------------------------------------------------------------

PRIMARY = [("WT_1", 10.0), ("A1", 20.0), ("B1", 5.0), ("C1", 12.0)]


def test_label_confirmation_refuses_wells_and_variants_of_the_run_together(tmp_path: Path, plate) -> None:
    confirm = _blocks(tmp_path / "c.xlsx", [("WT1", 10.0), ("5F-1", 20.0), ("B1-1", 5.0)])
    with pytest.raises(ValueError, match="cannot mix well and variant"):
        _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", PRIMARY), remeasure_report_xlsx=confirm)


def test_label_confirmation_sets_aside_what_the_run_does_not_hold(tmp_path: Path, plate) -> None:
    confirm = _blocks(
        tmp_path / "c.xlsx",
        [("WT1", 10.0), ("5F-1", 30.0), ("5F-2", 10.0), ("BLANK", 0.1), ("E5-1", 3.0), ("77Q-1", 4.0)],
    )
    result = _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", PRIMARY), remeasure_report_xlsx=confirm)
    assert result.variant_replicates["5F"] == pytest.approx([3.0, 1.0])
    assert result.exclusion_reason_counts == {"unrecognized_sample": 1, "unmapped_well": 1, "unmapped_variant": 1}


# ---------------------------------------------------------------------------
# Well-labelled replicated confirmation
# ---------------------------------------------------------------------------

def test_well_confirmation_refuses_this_runs_variant_label(tmp_path: Path, plate) -> None:
    confirm = _blocks(tmp_path / "c.xlsx", [("H12-1", 10.0), ("A1-1", 20.0), ("10L-1", 5.0)])
    with pytest.raises(ValueError, match="names variant 10L of this run"):
        _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", PRIMARY), remeasure_well_xlsx=confirm)


def test_well_confirmation_sets_aside_samples_of_no_membership(tmp_path: Path, plate) -> None:
    confirm = _blocks(tmp_path / "c.xlsx", [("H12-1", 10.0), ("A1-1", 20.0), ("BLANK", 0.1), ("77Q-1", 4.0), ("5", 1.0)])
    result = _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", PRIMARY), remeasure_well_xlsx=confirm)
    assert result.variant_replicates["5F"] == pytest.approx([2.0])
    assert result.exclusion_reason_counts == {"non_well_sample": 2, "numeric_calibration": 1}


# ---------------------------------------------------------------------------
# Rows the operator excluded
# ---------------------------------------------------------------------------

def test_an_operator_excluded_wt_row_leaves_the_reference(tmp_path: Path, plate) -> None:
    """Another experiment's WT under the same name has no signal; the operator names it."""
    report = _blocks(tmp_path / "r.xlsx", SHARED_RUN[:4] + [("WT4", 40.0)])
    excluded = [OperatorExclusion(source="round1_report_xlsx", block_index=5, row_index=23, sample_name="WT4", area=40.0)]
    result = _build(tmp_path, plate, round1_report_xlsx=report, excluded_rows=excluded)
    assert result.variant_replicates["5F"] == pytest.approx([2.0])
    assert result.exclusion_reason_counts == {"operator_excluded": 1}
    assert [row["sample_name"] for row in result.wt_rows] == ["WT1", "WT2", "WT3"]


def test_an_operator_exclusion_that_no_longer_matches_the_file_refuses(tmp_path: Path, plate) -> None:
    report = _blocks(tmp_path / "r.xlsx", SHARED_RUN[:4] + [("WT4", 40.0)])
    changed = [OperatorExclusion(source="round1_report_xlsx", block_index=5, row_index=23, sample_name="WT4", area=41.0)]
    with pytest.raises(ValueError, match="no longer holds the row that was excluded"):
        _build(tmp_path, plate, round1_report_xlsx=report, excluded_rows=changed)
    missing = [OperatorExclusion(source="round1_report_xlsx", block_index=9, row_index=43, sample_name="WT4", area=40.0)]
    with pytest.raises(ValueError, match="no input of this build holds"):
        _build(tmp_path, plate, round1_report_xlsx=report, excluded_rows=missing)


def test_an_operator_cannot_exclude_a_row_in_the_other_notation(tmp_path: Path, plate) -> None:
    """Rule 3 rows are the run's own measurements; unticking them would lose them silently."""
    report = _blocks(tmp_path / "r.xlsx", SHARED_RUN[:4] + [("11I", 9.0)])
    excluded = [OperatorExclusion(source="round1_report_xlsx", block_index=5, row_index=23, sample_name="11I", area=9.0)]
    with pytest.raises(ValueError, match="names variant 11I of this run"):
        _build(tmp_path, plate, round1_report_xlsx=report, excluded_rows=excluded)


ALL_ABOVE_WT = [("WT1", 10.0), ("1", 20.0), ("2", 15.0), ("3", 12.0)]


@pytest.mark.parametrize(
    ("field", "excluded_name"),
    [("numeric_report_xlsx", "BLANK"), ("numeric_report_xlsx", "WT2"), ("remeasure_numeric_xlsx", "WT2")],
)
def test_an_operator_exclusion_on_the_numeric_path_is_reported_once(tmp_path: Path, plate, field: str, excluded_name: str) -> None:
    rows = ALL_ABOVE_WT + [(excluded_name, 33.0)]
    report = _blocks(tmp_path / "n.xlsx", rows)
    excluded = [OperatorExclusion(source=field, block_index=5, row_index=23, sample_name=excluded_name, area=33.0)]
    sources: dict[str, Any] = {field: report}
    if field == "remeasure_numeric_xlsx":
        sources["numeric_report_xlsx"] = _blocks(tmp_path / "primary.xlsx", ALL_ABOVE_WT)
    result = _build(tmp_path, plate, excluded_rows=excluded, **sources)
    assert result.exclusion_reason_counts == {"operator_excluded": 1}
    assert [(row["source"], row["sample_name"], row["reason"]) for row in _manifest(result)["excluded_rows"]] == [
        (field, excluded_name, "operator_excluded"),
    ]
    assert sum(f"'{excluded_name}' (block 5, row 23)" in warning for warning in result.warnings) == 1


def test_well_confirmation_with_no_row_of_this_run_is_refused(tmp_path: Path, plate) -> None:
    confirm = _blocks(tmp_path / "c.xlsx", [("H12-1", 10.0), ("BLANK", 0.1), ("G7-1", 4.0)])
    with pytest.raises(ValueError, match="remeasure_well_xlsx has no row that belongs to this run"):
        _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", PRIMARY), remeasure_well_xlsx=confirm)


def test_reported_rows_are_the_rows_a_spreadsheet_shows(tmp_path: Path, plate) -> None:
    """The reader drops empty rows above the first cell; positions must not shift with it."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    row = 4  # three empty rows above, and an empty first column
    for name, area in SHARED_RUN:
        for offset, cells in enumerate((["Signal:", "FID1B"], ["Area", "Sample Name"], [area, name], ["Sum", area])):
            for column, value in enumerate(cells, start=2):
                ws.cell(row=row + offset, column=column, value=value)
        row += 5
    wb.save(tmp_path / "r.xlsx")
    result = _build(tmp_path, plate, round1_report_xlsx=tmp_path / "r.xlsx")
    (blank,) = result.excluded_rows
    assert (blank["sample_name"], blank["block_index"], blank["row_index"]) == ("BLANK", 5, 26)
    assert _cell(tmp_path / "r.xlsx", 26, 3) == "BLANK"

    flat = openpyxl.Workbook()
    sheet = flat.active
    assert sheet is not None
    for offset, cells in enumerate([["Sample Name", "Area"], ["A1", 2.0], ["BLANK", 0.1], ["B1", 0.5]]):
        for column, value in enumerate(cells, start=1):
            sheet.cell(row=3 + offset, column=column, value=value)
    flat.save(tmp_path / "gc.xlsx")
    result = _build(tmp_path, plate, gc_data_xlsx=tmp_path / "gc.xlsx")
    assert [(row["sample_name"], row["row_index"]) for row in result.excluded_rows] == [("BLANK", 5)]


@pytest.mark.parametrize("field", ["numeric_report_xlsx", "remeasure_numeric_xlsx"])
def test_an_operator_cannot_untick_this_runs_well_in_a_numeric_report(tmp_path: Path, plate, field: str) -> None:
    """Rule 3 comes before the operator's exclusions on the numeric path too."""
    report = _blocks(tmp_path / "n.xlsx", ALL_ABOVE_WT + [("C1", 9.0)])
    excluded = [OperatorExclusion(source=field, block_index=5, row_index=23, sample_name="C1", area=9.0)]
    sources: dict[str, Any] = {field: report}
    if field == "remeasure_numeric_xlsx":
        sources["numeric_report_xlsx"] = _blocks(tmp_path / "primary.xlsx", ALL_ABOVE_WT)
    with pytest.raises(ValueError, match="names well C01"):
        _build(tmp_path, plate, excluded_rows=excluded, **sources)


def test_a_confirmation_emptied_by_the_operator_says_so(tmp_path: Path, plate) -> None:
    confirm = _blocks(tmp_path / "c.xlsx", [("H12-1", 10.0), ("A1-1", 20.0)])
    excluded = [OperatorExclusion(source="remeasure_well_xlsx", block_index=2, row_index=8, sample_name="A1-1", area=20.0)]
    with pytest.raises(ValueError, match=r"every well of this run in it \(1 row\(s\)\) is in excluded_rows"):
        _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", PRIMARY), remeasure_well_xlsx=confirm, excluded_rows=excluded)


def test_rows_that_carry_only_formatting_do_not_shift_positions(tmp_path: Path, plate) -> None:
    """Styled but empty rows above the data: the reader skips them and the offset still holds."""
    from openpyxl.styles import Border, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    for row in (1, 2, 3):
        for column in (1, 2, 3, 4):
            cell = ws.cell(row=row, column=column)
            cell.fill = PatternFill("solid", fgColor="FFFF00")
            cell.border = Border(top=Side(style="thin"))
    row = 4
    for name, area in SHARED_RUN:
        for offset, cells in enumerate((["Signal:", "FID1B"], ["Area", "Sample Name"], [area, name], ["Sum", area])):
            for column, value in enumerate(cells, start=2):
                ws.cell(row=row + offset, column=column, value=value)
        row += 5
    wb.save(tmp_path / "styled.xlsx")
    result = _build(tmp_path, plate, round1_report_xlsx=tmp_path / "styled.xlsx")
    (blank,) = result.excluded_rows
    assert blank["row_index"] == 26
    assert _cell(tmp_path / "styled.xlsx", 26, 3) == "BLANK"


# ---------------------------------------------------------------------------
# A measurement on the plate WT well that no reference is taken from
# ---------------------------------------------------------------------------


def test_a_wt_well_measurement_in_a_block_report_is_reported_not_used(tmp_path: Path, plate) -> None:
    """The block report's reference is its named WT rows, so the H12 row is
    left out of the mean. It used to leave without a trace."""
    run = [*SHARED_RUN, ("H12", 9.0)]
    result = _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", run))
    assert dict(read_evolvepro_rows(tmp_path / "out.xlsx")) == pytest.approx({"5F": 2.0, "10L": 0.5, "11I": 1.2})
    assert "round1_report_xlsx excluded 1 row(s) (wt_well_not_reference): 'H12' (block 8, row 38)" in result.warnings
    assert result.exclusion_reason_counts == {"non_well_sample": 1, "wt_well_not_reference": 1}
    assert {
        "source": "round1_report_xlsx", "sample_name": "H12", "reason": "wt_well_not_reference",
        "row_index": 38, "block_index": 8, "area": 9.0, "kept": False,
    } in _manifest(result)["excluded_rows"]


def test_a_wt_well_row_in_a_flat_sheet_is_reported_not_used(tmp_path: Path, plate) -> None:
    sheet = _flat(tmp_path / "gc.xlsx", [("A1", 2.0), ("H12", 1.0), ("B1", 0.5)])
    result = _build(tmp_path, plate, gc_data_xlsx=sheet)
    assert dict(read_evolvepro_rows(tmp_path / "out.xlsx")) == pytest.approx({"5F": 2.0, "10L": 0.5})
    assert result.exclusion_reason_counts == {"wt_well_not_reference": 1}
    assert [(row["sample_name"], row["reason"], row["row_index"], row["block_index"]) for row in _manifest(result)["excluded_rows"]] == [
        ("H12", "wt_well_not_reference", 3, None),
    ]


def test_a_well_the_mapping_does_not_name_is_listed_with_its_position(tmp_path: Path, plate) -> None:
    run = [*SHARED_RUN, ("D1", 7.0), ("D1", 7.5)]
    result = _build(tmp_path, plate, round1_report_xlsx=_blocks(tmp_path / "r.xlsx", run))
    assert dict(read_evolvepro_rows(tmp_path / "out.xlsx")) == pytest.approx({"5F": 2.0, "10L": 0.5, "11I": 1.2})
    # One well, counted once, every read of it located.
    assert "1 measured well(s) absent from layout_xlsx: D01 (block 8, row 38; block 9, row 43)" in result.warnings
    assert result.exclusion_reason_counts["unmapped_well"] == 1
