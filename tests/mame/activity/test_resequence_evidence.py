"""Explicit clone correspondence preserves original activity and strict gating."""

import json
from pathlib import Path

import openpyxl
import pytest

from kuma_core.mame.activity.build_evolvepro_input import build_evolvepro_input
from kuma_core.mame.activity.evolvepro_xlsx import read_evolvepro_rows
from kuma_core.mame.activity.resequence_evidence import (
    CloneEvidenceLink,
    ResequenceEvidence,
    validate_resequence_links,
)


def _workbook(path: Path, *, repeated=False, identity="G2A", source_call="LOWDEPTH", reference="same-reference"):
    book = openpyxl.Workbook()
    final = book.active
    assert final is not None
    final.title = "Final"
    final.append(["well_id", "selected_plate", "mutant_id", "verdict", "is_fallback"])
    final.append(["C1" if repeated else "A1", "NB02" if repeated else "NB01", identity, "PASS" if repeated else source_call, not repeated])
    if not repeated:
        final.append(["B1", "NB01", "F3W", "PASS", False])
    for nb in ("NB01", "NB02"):
        sheet = book.create_sheet(nb)
        sheet.append(["well_id", "verdict"])
        sheet.append(["C1" if repeated else "A1", "PASS" if repeated else source_call])
        if not repeated:
            sheet.append(["B1", "PASS"])
    meta = book.create_sheet("__kuma_meta__")
    for row in (("reference_sha256", reference), ("coding_window", "0-9"), ("mode", "amplicon")):
        meta.append(row)
    book.save(path)
    return path


@pytest.fixture
def inputs(tmp_path):
    original = _workbook(tmp_path / "original.xlsx")
    repeated = _workbook(tmp_path / "repeated.xlsx", repeated=True)
    activity = tmp_path / "activity.csv"
    activity.write_text("well,value\nWT_1,1\nA1,1.7\nB1,2.3\n", encoding="utf-8")
    link = CloneEvidenceLink(relationship="same_clone", mutant_id="G2A", original_well="A1", original_plate="NB01", resequence_well="C1", resequence_plate="NB02")
    evidence = ResequenceEvidence(original_verdict_xlsx=str(original), verdict_xlsx=str(repeated), links=[link])
    return original, repeated, activity, evidence


def test_changed_layout_and_duplicate_nb_names_rescue_without_replacing_activity(tmp_path, inputs):
    original, repeated, activity, evidence = inputs
    before = original.read_bytes(), repeated.read_bytes()
    result = build_evolvepro_input(tmp_path / "out.xlsx", activity_path=activity, verdict_xlsx=original, resequence_evidence=evidence)
    assert dict(read_evolvepro_rows(result.output_path)) == pytest.approx({"2A": 1.7, "3W": 2.3})
    assert result.ngs_excluded == []
    assert result.exclusion_reason_counts == {}
    assert (original.read_bytes(), repeated.read_bytes()) == before
    assert result.manifest_path is not None
    manifest = json.loads(result.manifest_path.read_text())
    assert manifest["resequence_evidence"]["links"][0]["resequence_well"] == "C1"
    assert manifest["resequence_evidence"]["rescued_variants"] == ["2A"]


def test_without_explicit_link_original_gate_is_unchanged(tmp_path, inputs):
    original, _, activity, _ = inputs
    result = build_evolvepro_input(tmp_path / "out.xlsx", activity_path=activity, verdict_xlsx=original)
    assert dict(read_evolvepro_rows(result.output_path)) == {"3W": 2.3}
    assert result.ngs_excluded == ["2A"]


@pytest.mark.parametrize("change", ["identity", "reference", "contradiction", "duplicate", "wrong_plate", "different_clone"])
def test_incompatible_clone_evidence_is_rejected_without_output(tmp_path, inputs, change):
    original, repeated, activity, evidence = inputs
    if change == "identity":
        _workbook(repeated, repeated=True, identity="V2A")
    elif change == "reference":
        _workbook(repeated, repeated=True, reference="other-reference")
    elif change == "contradiction":
        _workbook(original, source_call="WRONG_AA")
    elif change == "duplicate":
        evidence = evidence.model_copy(update={"links": evidence.links * 2})
    elif change == "wrong_plate":
        evidence = evidence.model_copy(update={"links": [evidence.links[0].model_copy(update={"resequence_plate": "NB01"})]})
    else:
        with pytest.raises(ValueError):
            CloneEvidenceLink.model_validate({**evidence.links[0].model_dump(), "relationship": "new_clone"})
        return
    output = tmp_path / "out.xlsx"
    with pytest.raises(ValueError):
        build_evolvepro_input(output, activity_path=activity, verdict_xlsx=original, resequence_evidence=evidence)
    assert not output.exists()


def test_original_layout_identity_conflict_is_not_relaxed(tmp_path, inputs):
    original, _, activity, evidence = inputs
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Mutant", "Well Pos."])
    sheet.append(["G2A", "D1"])
    sheet.append(["F3W", "B1"])
    layout = tmp_path / "layout.xlsx"
    book.save(layout)
    activity.write_text("variant,value\nWT_1,1\n2A,1.7\n3W,2.3\n", encoding="utf-8")
    with pytest.raises(ValueError, match="conflicting layout"):
        build_evolvepro_input(tmp_path / "out.xlsx", activity_path=activity, layout_xlsx=layout, verdict_xlsx=original, resequence_evidence=evidence)


def test_final_activity_plate_combines_original_and_measured_resequence_clones(tmp_path, inputs):
    original, repeated, activity, evidence = inputs
    _workbook(original, source_call="WRONG_AA")
    activity.write_text("well,value\nWT_1,1\nD1,9.1\nB1,2.3\n", encoding="utf-8")
    measured = CloneEvidenceLink(
        relationship="measured_resequence_clone", mutant_id="G2A",
        activity_well="D1", resequence_well="C1", resequence_plate="NB02",
    )
    evidence = evidence.model_copy(update={"links": [measured]})
    before = original.read_bytes(), repeated.read_bytes()
    result = build_evolvepro_input(tmp_path / "out.xlsx", activity_path=activity, verdict_xlsx=original, resequence_evidence=evidence)
    assert dict(read_evolvepro_rows(result.output_path)) == pytest.approx({"2A": 9.1, "3W": 2.3})
    assert result.variant_replicates["2A"] == [9.1]
    assert result.n_ngs_excluded == 0
    assert (original.read_bytes(), repeated.read_bytes()) == before
    assert result.manifest_path is not None
    manifest = json.loads(result.manifest_path.read_text())
    assert manifest["resequence_evidence"]["links"][0]["relationship"] == "measured_resequence_clone"


def test_distinct_clone_activities_are_not_combined_by_mutant_id(tmp_path, inputs):
    original, _, activity, evidence = inputs
    activity.write_text("well,value\nWT_1,1\nA1,1.7\nD1,9.1\nB1,2.3\n", encoding="utf-8")
    measured = CloneEvidenceLink(
        relationship="measured_resequence_clone", mutant_id="G2A",
        activity_well="D1", resequence_well="C1", resequence_plate="NB02",
    )
    evidence = evidence.model_copy(update={"links": [measured]})
    with pytest.raises(ValueError, match="both original and replacement"):
        build_evolvepro_input(tmp_path / "out.xlsx", activity_path=activity, verdict_xlsx=original, resequence_evidence=evidence)
    assert not (tmp_path / "out.xlsx").exists()


def test_variant_only_activity_cannot_be_reinterpreted_as_a_new_clone(tmp_path, inputs):
    original, _, activity, evidence = inputs
    activity.write_text("variant,value\nWT_1,1\n2A,1.7\n3W,2.3\n", encoding="utf-8")
    measured = CloneEvidenceLink(
        relationship="measured_resequence_clone", mutant_id="G2A",
        activity_well="D1", resequence_well="C1", resequence_plate="NB02",
    )
    evidence = evidence.model_copy(update={"links": [measured]})
    with pytest.raises(ValueError, match="explicit well"):
        build_evolvepro_input(tmp_path / "out.xlsx", activity_path=activity, verdict_xlsx=original, resequence_evidence=evidence)


def test_measured_clone_relation_requires_an_activity_well():
    with pytest.raises(ValueError, match="explicit activity well"):
        CloneEvidenceLink(
            relationship="measured_resequence_clone", mutant_id="G2A",
            resequence_well="C1", resequence_plate="NB02",
        )


def test_measured_activity_well_may_be_absent_from_the_original_verdict(tmp_path, inputs):
    """Design, not an omission: a measured clone names a well the original never scored.

    `validate_resequence_links` deliberately leaves `source` unchecked on the
    measured branch. The original run did not place this clone, so there is no
    original identity to agree with, and docs/mame_resequencing.md calls the
    sample well A3 "the free well". The collision that would matter, an activity
    well the original assigned to a DIFFERENT variant, is refused one layer up
    by build_evolvepro_input against the active well->variant mapping, which is
    what the next two assertions pin.
    """
    original, _, activity, evidence = inputs
    measured = CloneEvidenceLink(
        relationship="measured_resequence_clone", mutant_id="G2A",
        activity_well="D1", resequence_well="C1", resequence_plate="NB02",
    )
    evidence = evidence.model_copy(update={"links": [measured]})
    # D1 is in no sheet of the original workbook, and the link still validates.
    assert validate_resequence_links(evidence, original) == {"2A": "D01"}

    activity.write_text("well,value\nWT_1,1\nD1,9.1\nB1,2.3\n", encoding="utf-8")
    result = build_evolvepro_input(
        tmp_path / "free.xlsx", activity_path=activity, verdict_xlsx=original,
        resequence_evidence=evidence,
    )
    assert dict(read_evolvepro_rows(result.output_path)) == pytest.approx({"2A": 9.1, "3W": 2.3})


def test_measured_activity_well_held_by_another_variant_is_refused(tmp_path, inputs):
    """The negative half of the rule above, with no layout sheet to appeal to."""
    original, _, activity, evidence = inputs
    # B1 is F3W in the original verdict, which is the derived mapping here.
    measured = CloneEvidenceLink(
        relationship="measured_resequence_clone", mutant_id="G2A",
        activity_well="B1", resequence_well="C1", resequence_plate="NB02",
    )
    evidence = evidence.model_copy(update={"links": [measured]})
    activity.write_text("well,value\nWT_1,1\nB1,2.3\n", encoding="utf-8")
    output = tmp_path / "collide.xlsx"
    with pytest.raises(ValueError, match="belongs to another variant"):
        build_evolvepro_input(output, activity_path=activity, verdict_xlsx=original, resequence_evidence=evidence)
    assert not output.exists()


def test_empty_plate_sheet_is_a_user_facing_error_not_an_indexerror(tmp_path, inputs):
    """A sheet with no rows must raise ValueError; the dispatcher only maps that."""
    original, repeated, _activity, evidence = inputs
    book = openpyxl.load_workbook(repeated)
    sheet = book["NB02"]
    book.remove(sheet)
    book.create_sheet("NB02")  # present, named, and completely empty
    book.save(repeated)
    with pytest.raises(ValueError, match="no well/verdict evidence"):
        validate_resequence_links(evidence, original)
