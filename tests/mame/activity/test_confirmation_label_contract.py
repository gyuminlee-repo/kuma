"""Synthetic boundary checks: source labels never silently become other wells."""
from pathlib import Path

import pytest
from openpyxl import Workbook

from kuma_core.mame.activity.build_evolvepro_input import (
    _confirmation, _strip_confirmation_rep_suffix, _read_long, _WT_RE,
)
from kuma_core.mame.activity.constants import WT_PATTERN
from kuma_core.mame.activity.detect_measurement_source import detect_measurement_source
from kuma_core.mame.activity.plate_layout_xlsx import _normalise_well


def report(path: Path, labels: list[str]) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for label in labels:
        sheet.append(["Signal:", "FID1B"])
        sheet.append(["Area", "Sample Name"])
        sheet.append([100, label])
        sheet.append(["Sum", 100])
        sheet.append([])
    book.save(path)
    return path


@pytest.mark.parametrize("label", ["WT", "wt", "Wt_1", "wt12", "WT_12"])
def test_wt_reader_and_classifier_share_contract(tmp_path: Path, label: str):
    assert _WT_RE is WT_PATTERN
    path = report(tmp_path / "report.xlsx", [label, "S11I-1"])
    assert _confirmation(path) == {"11I": [1.0]}
    detected = detect_measurement_source(path)
    assert detected.candidates == ["confirmationVariantLabels"]
    assert detected.evidence["n_sample_rows"] == 1
    assert detected.evidence["n_recognized_rows"] == 1
    assert detected.evidence["n_control_rows"] == 1
    assert "n_skipped_rows" not in detected.evidence
    assert detected.evidence["n_unique_labels"] == 1


@pytest.mark.parametrize("label", ["A0", "A00", "A13", "I1", "1-2", "A001", "A-1", "", "A1junk", "A1\n"])
def test_invalid_wells_are_refused_before_normalising(label: str):
    with pytest.raises(ValueError, match="96-well"):
        _normalise_well(label)


def test_every_96_coordinate_has_one_canonical_identity():
    wells = {_normalise_well(f"{row.lower()}{column}") for row in "ABCDEFGH" for column in range(1, 13)}
    assert len(wells) == 96
    assert wells == {f"{row}{column:02d}" for row in "ABCDEFGH" for column in range(1, 13)}


@pytest.mark.parametrize("label", ["S11I-0", "S11I--1", "S11I-1-2", "arbitrary-1", "S11I_Q12V-1", "A13-1", "1-2"])
def test_suffix_is_not_stripped_from_arbitrary_names(label: str):
    assert _strip_confirmation_rep_suffix(label) == label


@pytest.mark.parametrize("label, expected", [("S11I-1", "S11I"), ("11I-2", "11I"), ("A1-1", "A1"), ("h12-20", "h12")])
def test_explicit_canonical_replicate_suffix(label: str, expected: str):
    assert _strip_confirmation_rep_suffix(label) == expected


def test_well_confirmation_is_ambiguous_role_but_unambiguous_mapping(tmp_path: Path):
    path = report(tmp_path / "wells.xlsx", ["WT", "a1", "A01-2", "H12-1"])
    detected = detect_measurement_source(path)
    assert detected.candidates == ["rawReport", "confirmationWellLabels"]
    assert detected.ambiguous
    assert _confirmation(path, {"A01": "11I", "H12": "12V"}) == {"11I": [1.0, 1.0], "12V": [1.0]}
    with pytest.raises(ValueError, match="no unambiguous variant mapping"):
        _confirmation(path, {"A01": "11I"})


def test_confirmation_rejects_mixed_namespaces_and_empty_control_only(tmp_path: Path):
    with pytest.raises(ValueError, match="cannot mix"):
        _confirmation(report(tmp_path / "mixed.xlsx", ["WT", "A1", "S11I"]), {"A01": "11I"})
    with pytest.raises(ValueError, match="no measurement rows"):
        _confirmation(report(tmp_path / "empty.xlsx", ["WT"]))


def test_generic_reader_cannot_normalise_invalid_well_into_mapping(tmp_path: Path):
    path = tmp_path / "invalid.csv"
    path.write_text("well,value\nWT,1\nI1,2\n")
    with pytest.raises(ValueError, match="neither a canonical well nor variant"):
        _read_long(path, "raw", {"I01": "11I"}, {}, [])


def test_conflicting_layout_identity_is_not_last_row_wins(tmp_path: Path):
    from kuma_core.mame.activity.build_evolvepro_input import _layout_maps
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Mutant", "Well Pos."])
    sheet.append(["S11I", "A1"])
    sheet.append(["Q12V", "A01"])
    path = tmp_path / "conflict.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="well.*conflicting sample identities"):
        _layout_maps(path)


def test_template_bundle_and_preview_use_exact_synthetic_workbook():
    import json
    root = Path(__file__).resolve().parents[3]
    name = "17_mame_well_labeled_confirmation.xlsx"
    template = root / "templates" / name
    bundled = root / "src-tauri" / "samples" / "mame" / name
    assert template.read_bytes() == bundled.read_bytes()
    resources = json.loads((root / "src-tauri" / "tauri.conf.json").read_text())["bundle"]["resources"]
    assert resources[f"samples/mame/{name}"] == f"samples/mame/{name}"
    assert detect_measurement_source(template).candidates == ["rawReport", "confirmationWellLabels"]
    assert _confirmation(template, {"A01": "11I", "B01": "12V"}) == {"11I": [1.0, 1.0], "12V": [1.0]}


@pytest.mark.parametrize("first", ["WT", "WT_1", "blank", "unknown", "A11I"])
def test_layout_conflicts_are_rejected_before_skipping_control_or_unknown(first: str, tmp_path: Path):
    from kuma_core.mame.activity.build_evolvepro_input import _layout_maps
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Mutant", "Well Pos."])
    sheet.append([first, "A1"])
    sheet.append(["S11I", "A01"])
    path = tmp_path / "conflicting_layout.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="conflicting sample identities"):
        _layout_maps(path)


def test_replicated_raw_well_export_round_trips_through_gc_reader(tmp_path: Path):
    from kuma_core.mame.activity.build_evolvepro_input import build_evolvepro_input
    from kuma_core.mame.activity.evolvepro_xlsx import read_evolvepro_rows
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["well_id", "mutant_id", "verdict"])
    sheet.append(["A1", "S11I", "PASS"])
    verdict = tmp_path / "verdict.xlsx"
    book.save(verdict)
    raw = report(tmp_path / "raw.xlsx", ["WT", "a1", "A1-1"])
    gc = tmp_path / "gc.xlsx"
    first = build_evolvepro_input(tmp_path / "first.xlsx", round1_report_xlsx=raw, verdict_xlsx=verdict, gc_export_xlsx=gc)
    second = build_evolvepro_input(tmp_path / "second.xlsx", gc_data_xlsx=gc, verdict_xlsx=verdict)
    assert read_evolvepro_rows(first.output_path) == read_evolvepro_rows(second.output_path)
    assert first.variant_replicates == second.variant_replicates == {"11I": [1.0, 1.0]}


@pytest.mark.parametrize("sample", ["A11I", "A11I-1"])
def test_confirmation_full_identity_cannot_conflict_with_declared_reference(sample: str, tmp_path: Path):
    from kuma_core.mame.activity.build_evolvepro_input import build_evolvepro_input
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["well_id", "mutant_id", "verdict"])
    sheet.append(["A1", "S11I", "PASS"])
    verdict = tmp_path / "verdict.xlsx"
    book.save(verdict)
    primary = tmp_path / "primary.csv"
    primary.write_text("variant,value\n11I,1\n")
    confirmation = report(tmp_path / "confirmation.xlsx", ["WT", sample])
    output = tmp_path / "out.xlsx"
    with pytest.raises(ValueError, match="conflicts with declared identity"):
        build_evolvepro_input(output, activity_path=primary, activity_scale="relative_to_wt", remeasure_report_xlsx=confirmation, verdict_xlsx=verdict)
    assert not output.exists()
    for valid in ["S11I", "S11I-1", "11I", "11I-1"]:
        confirmation = report(tmp_path / "confirmation.xlsx", ["WT", valid])
        result = build_evolvepro_input(output, activity_path=primary, activity_scale="relative_to_wt", remeasure_report_xlsx=confirmation, verdict_xlsx=verdict)
        assert result.variant_replicates == {"11I": [1.0]}
