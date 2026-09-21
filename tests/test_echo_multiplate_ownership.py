import csv
from contextlib import closing
from dataclasses import asdict, replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from kuma_core.kuro.mutation import Mutation
from kuma_core.kuro.overlap import OverlapWindow
from kuma_core.kuro.plate_mapper import (
    build_echo_rows,
    deduplicate_reverse,
    echo_row_values,
    export_echo_mapping_csv,
    export_echo_mapping_xlsx,
    export_plate_excel,
    generate_plate_map,
)
from kuma_core.kuro.sdm_engine import SdmPrimerResult
from sidecar_kuro.handlers.export import (
    _build_echo_preview_rows,
    handle_export_echo_mapping_dry_run,
    handle_export_excel,
    handle_export_mapping,
)


def ownership_results(count: int) -> list[SdmPrimerResult]:
    results = []
    for i in range(count):
        reverse = "ACGT" * 5 + "ACGT"[(i // 32) % 4]
        results.append(SdmPrimerResult(
            mutation=Mutation(f"A{i + 2}V", "A", i + 2, "V", (i + 1) * 3, "GCT", "GTT"),
            forward_seq="GTT" + "ACGT" * 5, reverse_seq=reverse,
            forward_binding="ACGT" * 5, reverse_binding=reverse,
            overlap_window=OverlapWindow("ACGT" * 3, 0, 12, 12),
            tm_fwd=60, tm_rev=60, tm_overlap=40, tm_condition_met=True,
        ))
    return results


@pytest.mark.parametrize("count,quadrant,mapping_range", [
    (40, None, None), (97, None, None), (193, "A2", None),
    (97, None, ("C", "J")),
])
def test_echo_sources_match_order_workbook(
    tmp_path: Path, count: int, quadrant: str | None,
    mapping_range: tuple[str, str] | None,
) -> None:
    # Given generated mappings, including shared reverse groups across plates.
    results = ownership_results(count)
    fwd, rev = generate_plate_map(results, mapping_range=mapping_range)
    groups = deduplicate_reverse(results)
    order_path = tmp_path / "order.xlsx"
    csv_path = tmp_path / "echo.csv"
    xlsx_path = tmp_path / "echo.xlsx"

    # When the public export functions produce the physical order and worklists.
    export_plate_excel(fwd + rev, order_path, rev_groups=groups)
    rows = build_echo_rows(fwd, rev, groups, quadrant=quadrant, mapping_range=mapping_range)
    export_echo_mapping_csv(fwd, rev, csv_path, rev_groups=groups,
                            quadrant=quadrant, mapping_range=mapping_range)
    export_echo_mapping_xlsx(fwd, rev, xlsx_path, rev_groups=groups,
                             quadrant=quadrant, mapping_range=mapping_range)

    # Then each transfer addresses the ordered primer on that source plate.
    with closing(load_workbook(order_path)) as order:
        stock: dict[tuple[str, str], str] = {}
        for sheet in order:
            if " List" not in sheet.title:
                continue
            plate = sheet.title.split()[-1] if sheet.title[-1].isdigit() else "1"
            reverse = sheet.title.startswith("Rev")
            for well, name, *_ in sheet.iter_rows(min_row=2, values_only=True):
                row_index = "ABCDEFGH".index(str(well)[0]) * 2 + int(reverse)
                if quadrant is None and mapping_range:
                    row_index += "ABCDEFGHIJKLMNOP".index(mapping_range[0])
                col = int(str(well)[1:])
                if quadrant is not None:
                    # 96-head 는 한 칸 건너 열에만 닿는다. round 이름의 열 숫자가
                    # 열 offset 이다.
                    col = (col - 1) * 2 + 1 + (0 if quadrant == "A1" else 1)
                source_well = f"{'ABCDEFGHIJKLMNOP'[row_index]}{col}"
                stock[(f"Source [{plate}]", source_well)] = str(name)
        assert len(rows) == 2 * count
        for row in rows:
            assert stock.get((row["source_plate"], row["source_well"])) == row["source_well_name"], row

    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.reader(stream))[1:]
    assert csv_rows == [[str(value) for value in echo_row_values(row)] for row in rows]
    assert _build_echo_preview_rows(fwd, rev, 100, groups, quadrant=quadrant,
                                   mapping_range=mapping_range) == rows
    with closing(load_workbook(xlsx_path)) as workbook:
        assert list(workbook["Echo mapping file"].values)[1:] == [tuple(echo_row_values(row)) for row in rows]


def test_multiplate_layout_preserves_each_physical_plate(tmp_path: Path) -> None:
    # Given two plates whose A1 wells hold different primers.
    results = ownership_results(97)
    fwd, rev = generate_plate_map(results)
    path = tmp_path / "echo.xlsx"

    # When exporting the human-readable loading instructions.
    export_echo_mapping_xlsx(fwd, rev, path, rev_groups=deduplicate_reverse(results))

    # Then each worklist source and destination has a separately labelled grid.
    with closing(load_workbook(path)) as workbook:
        assert "layout 2" in workbook.sheetnames
        rows = list(workbook["Echo mapping file"].values)[1:]
        for index, title in enumerate(("layout", "layout 2"), 1):
            sheet = workbook[title]
            assert sheet.cell(2, 1).value == f"Source [{index}]"
            assert sheet.cell(22, 1).value == f"Destination [{index}]"
            for source, name, well, dest, mutation, dest_well, _ in rows:
                assert isinstance(well, str)
                assert isinstance(dest_well, str)
                if source == f"Source [{index}]":
                    assert sheet.cell(5 + "ABCDEFGHIJKLMNOP".index(well[0]), 2 + int(well[1:])).value == name
                if dest == f"Destination [{index}]":
                    assert sheet.cell(25 + "ABCDEFGH".index(dest_well[0]), 2 + int(dest_well[1:])).value == mutation
            usage = list(sheet.iter_rows(min_row=36, max_row=sheet.max_row - 1, values_only=True))
            counts = [row[2] for row in usage]
            assert all(isinstance(count, int) for count in counts)
            assert sum(int(str(count)) for count in counts) == (96 if index == 1 else 1)


@pytest.mark.parametrize("count,mapping_range,well_order", [
    (5, ("C", "J"), "column"),
    (48, ("C", "J"), "column"),
    (49, ("C", "J"), "column"),
    (97, ("C", "J"), "column"),
    (49, ("C", "J"), "row"),
    (25, ("C", "F"), "column"),
    (13, ("J", "K"), "column"),
    (5, None, "column"),
    (97, None, "column"),
    (97, None, "row"),
])
def test_unique_reverse_range_exports_match_physical_stock(
    tmp_path: Path, count: int, mapping_range: tuple[str, str] | None,
    well_order: str,
) -> None:
    # Given enough unique reverse primers to cross the reduced-row boundary.
    results = [
        replace(result, reverse_seq="ACGT" + format(i, "08b").translate(str.maketrans("01", "AC")))
        for i, result in enumerate(ownership_results(count))
    ]
    fwd, rev = generate_plate_map(results, well_order=well_order, mapping_range=mapping_range)
    payload = {"mappings": [asdict(m) for m in fwd + rev], "dedup_info": deduplicate_reverse(results)}
    placement = {"mapping_range": {"row_start": mapping_range[0], "row_end": mapping_range[1]}} if mapping_range else {}
    order_path = tmp_path / "order.xlsx"

    # When the real handlers export ordered stock, preview and both worklists.
    assert handle_export_excel({**payload, "filepath": str(order_path)})["success"]
    rows = handle_export_echo_mapping_dry_run({**payload, **placement})["rows"]
    for extension in ("csv", "xlsx"):
        assert handle_export_mapping({
            **payload, **placement, "format": "echo",
            "filepath": str(tmp_path / f"echo.{extension}"),
        })["success"]

    # Then every transfer names the actual ordered stock and no plate overflows.
    capacity = 96 if mapping_range is None else (ord(mapping_range[1]) - ord(mapping_range[0]) + 1) * 6
    plate_count = (count + capacity - 1) // capacity
    start = 0 if mapping_range is None else ord(mapping_range[0]) - ord("A")
    stock: dict[tuple[str, str], str] = {}
    with closing(load_workbook(order_path)) as workbook:
        for plate in range(1, plate_count + 1):
            tag = f" {plate}" if plate_count > 1 else ""
            for direction in ("Fwd", "Rev"):
                ordered = list(workbook[f"{direction} List{tag}"].values)[1:]
                assert len(ordered) == min(capacity, count - (plate - 1) * capacity)
                for well, name, *_ in ordered:
                    assert isinstance(well, str)
                    row = start + 2 * (ord(well[0]) - ord("A")) + (direction == "Rev")
                    address = (f"Source [{plate}]", f"{chr(ord('A') + row)}{well[1:]}")
                    assert address not in stock
                    stock[address] = str(name)
    assert len(rows) == count * 2
    assert {(row["source_plate"], row["source_well"]) for row in rows} == set(stock)
    for row in rows:
        assert stock[(row["source_plate"], row["source_well"])] == row["source_well_name"]
    if well_order == "column" and mapping_range == ("C", "J"):
        assert rows[count + 4]["source_well"] == "D2"
    if well_order == "column" and mapping_range is None:
        assert rows[count + 4]["source_well"] == "J1"
    with (tmp_path / "echo.csv").open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.reader(stream))[1:] == [[str(v) for v in echo_row_values(row)] for row in rows]
    with closing(load_workbook(tmp_path / "echo.xlsx")) as workbook:
        assert list(workbook["Echo mapping file"].values)[1:] == [tuple(echo_row_values(row)) for row in rows]
        for row in rows:
            plate = int(row["source_plate"].split("[")[1][:-1])
            sheet = workbook["layout" if plate == 1 else f"layout {plate}"]
            well = row["source_well"]
            assert sheet.cell(5 + ord(well[0]) - ord("A"), 2 + int(well[1:])).value == row["source_well_name"]


def test_echo_range_does_not_repack_default_physical_forward_stock() -> None:
    fwd, rev = generate_plate_map(ownership_results(5))
    assert fwd[4].well == "E1"
    with pytest.raises(ValueError, match="source row E would wrap"):
        build_echo_rows(fwd, rev, deduplicate_reverse(ownership_results(5)), mapping_range=("C", "J"))
