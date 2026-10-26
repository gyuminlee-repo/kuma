"""Inert regressions for audit file/metadata boundaries.

Compile only the exact functions under test when their enclosing module imports
scientific engines. Renderers/parsers are stubs: these tests do not exercise
sequence design, analysis, model selection, instruments, or live inputs.
"""
from __future__ import annotations

import ast
import contextlib
import csv
import datetime
import importlib.util
import io
import json
import logging
import math
import os
from pathlib import Path
import queue
import re
import shutil
import subprocess
import sys
import threading
from typing import NamedTuple
from types import SimpleNamespace
import zipfile

import openpyxl
import pytest

from kuma_core.shared import atomic_write

ROOT = Path(__file__).resolve().parents[2]


def boundary(relative, names, **extra):
    path = ROOT / relative
    tree = ast.parse(path.read_text(encoding="utf-8"))
    nodes = []
    for node in tree.body:
        if isinstance(node, ast.ImportFrom) and node.module == "__future__":
            nodes.append(node)
        elif isinstance(node, (ast.FunctionDef, ast.ClassDef)) and node.name in names:
            nodes.append(node)
        elif isinstance(node, (ast.Assign, ast.AnnAssign)):
            targets = node.targets if isinstance(node, ast.Assign) else [node.target]
            if any(isinstance(t, ast.Name) and t.id in names for t in targets):
                nodes.append(node)
    env: dict[str, object] = dict(Path=Path, json=json, csv=csv, io=io, math=math, re=re,
               datetime=datetime, zipfile=zipfile, contextlib=contextlib,
               threading=threading, queue=queue, shutil=shutil,
               subprocess=subprocess, logger=logging.getLogger(__name__),
               atomic_write_text=atomic_write.atomic_write_text)
    if hasattr(atomic_write, "atomic_output_path"):
        env["atomic_output_path"] = atomic_write.atomic_output_path
    env.update(extra)
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(path), "exec"), env)
    return SimpleNamespace(**env)


def load_metadata_module(name):
    path = ROOT / "kuma_core/mame/ingest" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(f"audit_boundary_{name}", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def pdf_boundary():
    return boundary("kuma_core/mame/report/pdf_export.py",
                    {"export_pdf", "_html_fallback_path", "_validate_pdf"},
                    render_html=lambda _: "<p>generic report</p>")


@pytest.mark.parametrize("failure", ["error", "timeout", "invalid"])
def test_core01_pdf_failure_preserves_existing_file(tmp_path, monkeypatch, failure):
    output = tmp_path / "report.pdf"
    output.write_bytes(b"EXISTING REPORT")
    monkeypatch.setattr(shutil, "which", lambda _: "generic-converter")

    def convert(args, **kwargs):
        Path(args[-1]).write_bytes(b"PARTIAL")
        if failure == "timeout":
            raise subprocess.TimeoutExpired(args, 60)
        return SimpleNamespace(returncode=1 if failure == "error" else 0, stderr=b"failed")

    monkeypatch.setattr(subprocess, "run", convert)
    result = pdf_boundary().export_pdf(object(), output)
    assert result["error"]
    assert output.read_bytes() == b"EXISTING REPORT"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["report.pdf"]


def test_core01_pdf_success_publishes_validated_stage(tmp_path, monkeypatch):
    output = tmp_path / "report.pdf"
    output.write_bytes(b"OLD")
    monkeypatch.setattr(shutil, "which", lambda _: "generic-converter")
    rendered = b"%PDF-1.7\ngeneric inert fixture\n%%EOF\n"

    def convert(args, **kwargs):
        assert Path(args[-1]) != output
        assert Path(args[-1]).parent == output.parent
        Path(args[-1]).write_bytes(rendered)
        assert output.read_bytes() == b"OLD"
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(subprocess, "run", convert)
    result = pdf_boundary().export_pdf(object(), output)
    assert result["error"] is None
    assert output.read_bytes() == rendered
    assert sorted(p.name for p in tmp_path.iterdir()) == ["report.pdf"]


def test_core01_pdf_fallback_reexport_replaces_the_previous_html(tmp_path, monkeypatch):
    """The shipped app has no weasyprint, so every PDF export lands here.

    Exporting twice to one name must succeed twice and leave the second report,
    and must not need hard links, which exFAT does not have.
    """
    output = tmp_path / "report.pdf"
    fallback = output.with_suffix(".html")
    monkeypatch.setattr(shutil, "which", lambda _: None)
    monkeypatch.setattr(importlib.util, "find_spec", lambda _: None)
    monkeypatch.setattr(atomic_write.os, "link", lambda *_, **__: (_ for _ in ()).throw(
        OSError("hard links are not supported")), raising=False)
    for body in ["<p>first</p>", "<p>second</p>"]:
        env = boundary("kuma_core/mame/report/pdf_export.py",
                       {"export_pdf", "_html_fallback_path", "_validate_pdf"},
                       render_html=lambda _, body=body: body)
        result = env.export_pdf(object(), output)
        assert result["error"] is None
        assert result["output_path"] == str(fallback)
        assert fallback.read_text(encoding="utf-8") == body
    assert list(tmp_path.iterdir()) == [fallback]


@pytest.mark.parametrize("handler", ["handle_save_json", "handle_save_workspace"])
def test_core01_json_serialization_failure_preserves_file(tmp_path, handler):
    output = tmp_path / "state.json"
    output.write_text('{"old": true}')
    env = boundary("python-core/sidecar_kuro/handlers/export.py", {handler},
                   SaveJsonParams=SimpleNamespace, SaveWorkspaceParams=SimpleNamespace,
                   _validate_output_path=lambda path, **_: Path(path))
    with pytest.raises(TypeError):
        getattr(env, handler)({"filepath": str(output), "data": {"bad": object()}})
    assert output.read_text(encoding="utf-8") == '{"old": true}'
    assert list(tmp_path.iterdir()) == [output]


@pytest.mark.parametrize("handler", ["handle_save_json", "handle_save_workspace"])
def test_core01_json_publish_failure_preserves_file(tmp_path, monkeypatch, handler):
    output = tmp_path / "state.json"
    output.write_text('{"old": true}')
    env = boundary("python-core/sidecar_kuro/handlers/export.py", {handler},
                   SaveJsonParams=SimpleNamespace, SaveWorkspaceParams=SimpleNamespace,
                   FileExportResultModel=lambda **kw: SimpleNamespace(to_rpc_dict=lambda: kw),
                   _validate_output_path=lambda path, **_: Path(path))
    monkeypatch.setattr(atomic_write, "_replace_with_retry", lambda *_: (_ for _ in ()).throw(OSError("disk failure")))
    with pytest.raises(OSError, match="disk failure"):
        getattr(env, handler)({"filepath": str(output), "data": {"new": True}})
    assert output.read_text(encoding="utf-8") == '{"old": true}'
    assert list(tmp_path.iterdir()) == [output]


@pytest.mark.parametrize("filename,function", [
    ("excel_writer.py", "_write_kuma_meta_sheet"),
    ("janus_mapping.py", "_write_janus_kuma_meta_sheet"),
])
def test_core02_metadata_roundtrips_as_literal_text(tmp_path, filename, function):
    env = boundary(f"kuma_core/mame/export/{filename}", {function},
                   _style_header=lambda _: None, _finalize=lambda *a, **kw: None,
                   VerdictClass=[])
    meta = SimpleNamespace(**dict.fromkeys(
        ["instrument", "position", "flow_cell_id", "sample_id", "kit", "started", "raw_run_dir"], "=1+1"),
        basecalling_enabled=False)
    wb = openpyxl.Workbook()
    getattr(env, function)(wb, meta, "generic")
    path = tmp_path / "metadata.xlsx"
    wb.save(path)
    wb.close()
    with contextlib.closing(openpyxl.load_workbook(path, data_only=False)) as saved:
        for row in saved["__kuma_meta__"].iter_rows():
            if row[1].value == "=1+1":
                assert row[1].data_type == "s"
        assert saved["__kuma_meta__"].max_row > 8


@pytest.mark.parametrize("payload", [[], "bad", {"protocol_run_info": "bad"},
    {"acquisitions": 12}, {"acquisitions": [{"acquisition_run_info": "bad"}]},
    {"acquisitions": [{"acquisition_run_info": {"bream_info": "bad"}}]},
    {"acquisitions": [{"acquisition_run_info": {"bream_info": {"mux_scan_results": 12}}}]},
    {"acquisitions": [{"acquisition_run_info": {"bream_info": {"mux_scan_results": [{"counts": "bad"}]}}}]},
])
def test_core03_malformed_optional_json_is_absent(tmp_path, payload):
    (tmp_path / "report_generic.json").write_text(json.dumps(payload))
    env = load_metadata_module("flow_cell")
    result = env.read_flow_cell_history(tmp_path)
    assert result.scans == []
    json.dumps(env.serialise_history(result), allow_nan=False)


@pytest.mark.parametrize("stamp", [
    float("nan"), float("inf"), float("-inf"), True,
    "nan", "inf", "-inf", "abc", "", "True", -5, -1.5, "-5", [120], {"s": 120}])
def test_core03_invalid_timestamp_is_not_a_scan(tmp_path, stamp):
    payload = {"acquisitions": [{"acquisition_run_info": {"bream_info": {
        "mux_scan_results": [{"counts": {"single_pore": 7}, "mux_scan_timestamp": stamp}]
    }}}]}
    (tmp_path / "report_generic.json").write_text(json.dumps(payload))
    assert load_metadata_module("flow_cell").read_flow_cell_history(tmp_path).scans == []


@pytest.mark.parametrize("stamp,seconds", [("120", 120), ("1.5", 1), ("0", 0), (120, 120), (1.5, 1)])
def test_core03_numeric_string_timestamp_is_a_scan(tmp_path, stamp, seconds):
    payload = {"acquisitions": [{"acquisition_run_info": {"bream_info": {
        "mux_scan_results": [{"counts": {"single_pore": 7}, "mux_scan_timestamp": stamp}]
    }}}]}
    (tmp_path / "report_generic.json").write_text(json.dumps(payload))
    scans = load_metadata_module("flow_cell").read_flow_cell_history(tmp_path).scans
    assert [(s.at_seconds, s.single_pore) for s in scans] == [(seconds, 7)]


def test_core03_minknow_shaped_report_keeps_every_mux_scan():
    """MinKNOW writes ``mux_scan_timestamp`` as a string of digits.

    The fixture is synthetic but has the same shape as a real MinKNOW report
    reduced to the two paths MAME reads: the same keys, the same value types
    (digit-string stamps, integer pore counts), a first acquisition without
    ``bream_info`` and eight mux scans. Identity, pore counts and stamps are
    invented values.
    """
    run_dir = ROOT / "tests/mame/fixtures/minknow_report_synthetic"
    env = load_metadata_module("flow_cell")
    result = env.read_flow_cell_history(run_dir)
    assert (result.flow_cell_id, result.product_code, result.channel_count) == ("FAX00001", "FLO-MIN114", 512)
    assert [(s.at_seconds, s.single_pore) for s in result.scans] == [
        (120, 400), (5720, 371), (11320, 338), (16920, 297),
        (22520, 312), (28120, 254), (33720, 268), (39320, 200)]
    assert (result.pore_start, result.pore_end) == (400, 200)
    json.dumps(env.serialise_history(result), allow_nan=False)


def test_p04_ledger_entry_written_under_the_former_identity_is_this_run(tmp_path):
    """A raw run used to be recorded under its demux output directory.

    Re-analysing that run now names it by the run folder. The older entry is
    the same campaign, so it must neither read as an earlier use of the cell
    nor survive beside the new entry.
    """
    env = load_metadata_module("flow_cell")
    ledger = tmp_path / env.FLOW_CELL_LEDGER_NAME
    ledger.write_text(json.dumps([
        {"flow_cell_id": "CELL1", "run_dir": "elsewhere/run", "started": "earlier"},
        {"flow_cell_id": "CELL1", "run_dir": "project/demux", "started": "this run"},
    ]))
    entries = env.read_ledger(tmp_path)
    previous = env.find_previous_use(entries, "CELL1", "source/run", former_run_dir="project/demux")
    assert previous is not None and previous["run_dir"] == "elsewhere/run"
    assert env.find_previous_use(entries[1:], "CELL1", "source/run", former_run_dir="project/demux") is None
    # Without the former identity the same entry is somebody else's run.
    assert env.find_previous_use(entries[1:], "CELL1", "source/run") == entries[1]
    env.record_use(tmp_path, env.FlowCellHistory(flow_cell_id="CELL1"), "source/run", "now",
                   former_run_dir="project/demux")
    assert [e["run_dir"] for e in env.read_ledger(tmp_path)] == ["elsewhere/run", "source/run"]


@pytest.mark.parametrize("payload", [
    {"acquisitions": 12},
    {"acquisitions": [{"read_length_histogram": 12}]},
    {"acquisitions": [{"read_length_histogram": True, "qscore_histograms": 3.5}]},
    {"acquisitions": [{"qscore_histograms": 12}]},
])
def test_core03_malformed_read_length_containers_are_absent(tmp_path, payload):
    from kuma_core.mame.ingest.read_length import read_read_length_qc
    (tmp_path / "report_generic.json").write_text(json.dumps(payload))
    result = read_read_length_qc(tmp_path)
    assert result.histograms is None
    assert result.qscore_histograms is None


def zip_boundary():
    names = {"_load_from_zip", "load_structure_file", "StructureFileError", "MAX_BYTES",
             "MAX_ZIP_MEMBERS", "CIF_SUFFIXES", "PDB_SUFFIXES", "ZIP_SUFFIXES", "SUPPORTED_SUFFIXES",
             "_validate_zip_limits", "_read_zip_member", "_summary_ranking_score", "_MODEL_INDEX_RE",
             "_summary_candidates", "ModelCandidate", "LoadedStructure"}
    from dataclasses import dataclass
    return boundary("kuma_core/kuro/structure_file.py", names, dataclass=dataclass,
                    _parse_cif_ca=lambda _: pytest.fail("domain parser must not run"))


@pytest.mark.parametrize("entries", [
    [("generic.cif", " " * 1025)],
    [("a.cif", " " * 600), ("b.cif", " " * 600)],
    [("generic_model_0.cif", "x"), ("generic_summary_confidences_0.json", " " * 1025)],
])
def test_core04_zip_expansion_is_bounded_before_parser(tmp_path, entries):
    env = zip_boundary()
    env.load_structure_file.__globals__["MAX_BYTES"] = 1024
    path = tmp_path / "generic.zip"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    with pytest.raises(env.StructureFileError, match="limit"):
        env.load_structure_file(path)


def test_core04_members_the_loader_never_opens_do_not_count(tmp_path):
    """An AlphaFold Server archive carries per-model ``full_data`` JSON that is
    far larger than the models and that the loader never decompresses."""
    env = zip_boundary()
    scope = env.load_structure_file.__globals__
    scope["MAX_BYTES"] = 1024
    scope["_parse_cif_ca"] = lambda _: ([None, (0.0, 0.0, 0.0)], "A", 50.0)
    path = tmp_path / "generic.zip"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("generic_model_0.cif", "x")
        archive.writestr("generic_summary_confidences_0.json", '{"ranking_score": 0.5}')
        archive.writestr("generic_full_data_0.json", " " * 4096)
    loaded = env.load_structure_file(path)
    assert loaded.source_name == "generic_model_0.cif"
    assert loaded.candidates[0].ranking_score == 0.5


def test_core05_close_stops_producer_even_after_queue_drain():
    waiting = threading.Event()
    release = threading.Event()
    made_queues = []

    class ControlledQueue(queue.Queue):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            made_queues.append(self)

        def put(self, item, block=True, timeout=None):
            if item == (True, 1):
                waiting.set()
                assert release.wait(3)
            return super().put(item, block, timeout)

    env = boundary("kuma_core/mame/ingest/combinatorial_demux.py", {"_prefetch"},
                   queue=SimpleNamespace(Queue=ControlledQueue, Empty=queue.Empty, Full=queue.Full))
    before = {t.ident for t in threading.enumerate()}
    iterator = env._prefetch(iter([0, 1]), 1)
    assert next(iterator) == 0
    assert waiting.wait(3)
    closer = threading.Thread(target=iterator.close)
    closer.start()
    # Producer is released after close starts. A bounded join inside close must
    # not depend on an uninterruptible queue.put for its completion marker.
    release.set()
    closer.join(3)
    workers = [t for t in threading.enumerate() if t.ident not in before and t.name == "fastq-prefetch"]
    for worker in workers:
        worker.join(.5)
    try:
        assert not closer.is_alive()
        assert not any(t.is_alive() for t in workers)
    finally:
        # Clean up even against the old implementation, which leaks at _DONE.
        for q in made_queues:
            with contextlib.suppress(queue.Empty):
                q.get_nowait()
        for worker in workers:
            worker.join(1)


def test_core05_producer_exception_and_order_are_preserved():
    env = boundary("kuma_core/mame/ingest/combinatorial_demux.py", {"_prefetch"})
    assert list(env._prefetch(iter(range(5)), 1)) == list(range(5))
    def source():
        yield 1
        raise ValueError("generic read failure")
    iterator = env._prefetch(source(), 1)
    assert next(iterator) == 1
    with pytest.raises(ValueError, match="generic read failure"):
        next(iterator)


@pytest.mark.parametrize("linked", ["directory", "summary", "sample_sheet", "input"])
def test_core06_discovery_does_not_follow_metadata_symlinks(tmp_path, linked):
    scope = tmp_path / "scope"
    scope.mkdir()
    source = scope / "input"
    source.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "final_summary_other.txt").write_text("sample_id = elsewhere\n")
    (outside / "sample_sheet_other.csv").write_text("kit,external\n")
    if linked == "directory":
        (scope / "linked").symlink_to(outside, target_is_directory=True)
    elif linked == "summary":
        (source / "final_summary_link.txt").symlink_to(outside / "final_summary_other.txt")
    elif linked == "sample_sheet":
        (source / "final_summary_own.txt").write_text("sample_id = own\n")
        (source / "sample_sheet_link.csv").symlink_to(outside / "sample_sheet_other.csv")
    else:
        source = scope / "linked_input"
        source.symlink_to(outside, target_is_directory=True)
    env = load_metadata_module("run_meta")
    setattr(env, "_MAX_LEVELS_UP", 1)
    result = env.discover_run_meta(source)
    if linked == "sample_sheet":
        assert result.sample_id == "own"
        assert result.kit is None
    else:
        assert result is None


def throughput_boundary():
    return boundary("kuma_core/mame/health.py", {
        "_parse_throughput", "_read_csv_rows", "_first_col", "_TIME_COLS",
        "_READS_PER_SEC_COLS", "_CUMULATIVE_READS_COLS"})


@pytest.mark.parametrize("bad", ["nan", "inf", "-inf"])
@pytest.mark.parametrize("column", ["time", "rate", "cumulative"])
def test_core07_nonfinite_optional_rows_never_reach_json(tmp_path, bad, column):
    env = throughput_boundary()
    time_col = env._TIME_COLS[0]
    metric_col = env._CUMULATIVE_READS_COLS[0] if column == "cumulative" else env._READS_PER_SEC_COLS[0]
    middle = f"{bad},10" if column == "time" else f"1,{bad}"
    (tmp_path / "throughput_generic.csv").write_text(f"{time_col},{metric_col}\n0,0\n{middle}\n2,120\n")
    result = env._parse_throughput(tmp_path)
    json.dumps(result, allow_nan=False)
    assert result
    if column == "cumulative":
        assert result == [{"time_h": .0333, "reads_per_sec": 1.0}]


def test_p03_preview_only_consumes_requested_rows_and_closes(monkeypatch):
    seen = []
    closed = []
    def rows(**kwargs):
        for i in range(10001):
            seen.append(i)
            yield (f"item-{i}",)
    wb = SimpleNamespace(sheetnames=["Data"], close=lambda: closed.append(True))
    class Workbook:
        sheetnames = wb.sheetnames
        def __getitem__(self, name):
            return SimpleNamespace(iter_rows=rows)
        def close(self):
            wb.close()
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **kw: Workbook())
    env = boundary("python-core/sidecar_kuro/handlers/misc.py", {"_preview_xlsx"})
    result = env._preview_xlsx("generic.xlsx", None, 3)
    assert len(result["rows"]) == 3
    assert seen == [0, 1, 2, 3]
    assert closed == [True]


def test_p03_preview_closes_workbook_on_read_error(monkeypatch):
    closed = []
    class Workbook:
        sheetnames = ["Data"]
        def __getitem__(self, name):
            raise OSError("bad worksheet")
        def close(self):
            closed.append(True)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **kw: Workbook())
    env = boundary("python-core/sidecar_kuro/handlers/misc.py", {"_preview_xlsx"})
    with pytest.raises(OSError, match="bad worksheet"):
        env._preview_xlsx("generic.xlsx", None, 3)
    assert closed == [True]


@pytest.mark.parametrize("raw", [True, False])
def test_p04_flow_metadata_uses_original_source_not_rebound_output(tmp_path, raw):
    source = tmp_path / "source" / "run"
    output = tmp_path / "elsewhere" / "intermediate"
    source.mkdir(parents=True)
    output.mkdir(parents=True)
    (source / "report_generic.json").write_text(json.dumps({"protocol_run_info": {"flow_cell": {"flow_cell_id": "original"}}}))
    (output / "report_generic.json").write_text(json.dumps({"protocol_run_info": {"flow_cell": {"flow_cell_id": "intermediate"}}}))
    tree = ast.parse((ROOT / "python-core/sidecar_mame/handlers/analyze.py").read_text(encoding="utf-8"))
    assignment = next(n for n in ast.walk(tree) if isinstance(n, ast.Assign)
                      and any(isinstance(t, ast.Name) and t.id == "_flow_cell" for t in n.targets))
    env = {"read_flow_cell_history": load_metadata_module("flow_cell").read_flow_cell_history,
           "original_run_dir": source, "input_dir": output, "is_raw": raw}
    exec(compile(ast.Module(body=[assignment], type_ignores=[]), "metadata_call", "exec"), env)
    assert env["_flow_cell"].flow_cell_id == ("original" if raw else "intermediate")


def long_csv_boundary():
    import pandas as pd
    from typing import cast
    path = ROOT / "kuma_core/mame/activity/models.py"
    spec = importlib.util.spec_from_file_location("audit_activity_models", path)
    assert spec is not None and spec.loader is not None
    models = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = models
    spec.loader.exec_module(models)
    env = {name: getattr(models, name) for name in [
        "ActivityRecord", "ActivityTable", "DropReason", "DroppedRow", "PlateConfig", "PlateMeta", "WtReplicateRecord"]}
    from decimal import Decimal, InvalidOperation
    return boundary("kuma_core/mame/activity/ingest_long_csv.py", {
        "ingest_long_csv", "_normalise_well", "_is_valid_well", "_DETAIL_MAX_CHARS",
        "WELL_RE_96", "WELL_RE_384", "WELL_RE_RAW", "_WT_SUFFIX_RE", "WELL_COL_ALIASES", "VALUE_COL_ALIASES"},
        pd=pd, cast=cast, WT_PATTERN=re.compile(r"^WT_?\d+$"),
        Decimal=Decimal, InvalidOperation=InvalidOperation, **env)


def test_new_io01_fractional_replicate_ids_are_rejected_not_merged(tmp_path):
    path = tmp_path / "generic.csv"
    path.write_text("plate_id,well_id,value,replicate_idx\nP1,A01,10,1.2\nP1,A01,30,1.8\nP1,A02,20,2.0\n")
    result = long_csv_boundary().ingest_long_csv(path, {"P1": []})
    assert [r.replicate_idx for r in result.records] == [2]
    assert [r.row_index for r in result.dropped_rows] == [0, 1]
    assert {r.reason for r in result.dropped_rows} == {"replicate_idx_unparseable"}


def test_new_io02_blank_plate_id_is_not_the_literal_nan(tmp_path):
    path = tmp_path / "generic.csv"
    path.write_text("plate_id,well_id,value\n,A01,10\n  ,WT_1,5\nP1,A02,20\n")
    result = long_csv_boundary().ingest_long_csv(path, {"P1": ["A01"]})
    assert [r.plate_id for r in result.records] == ["P1"]
    assert not result.wt_records
    assert [r.row_index for r in result.dropped_rows] == [0, 1]
    assert {r.reason for r in result.dropped_rows} == {"plate_id_missing"}


def read_length_boundary():
    from dataclasses import dataclass, field
    return boundary("kuma_core/mame/ingest/read_length.py", {
        "_as_int", "_as_float", "_as_nonnegative_count", "_read_buckets", "_read_qscore", "Buckets", "QScoreHistogram", "QScoreSeries"},
        dataclass=dataclass, field=field, __name__=__name__)


def test_new_io03_malformed_qscore_series_is_omitted_without_shifting():
    result = read_length_boundary()._read_qscore({
        "bucket_ranges": [{"start": i, "end": i + 1} for i in range(3)],
        "histogram_data": [{"bucket_values": [10, "bad", 30]}, {"bucket_values": [40, 50, 60]}]})
    assert result.starts == [0.0, 1.0, 2.0]
    assert result.ends == [1.0, 2.0, 3.0]
    assert [series.values for series in result.series] == [[40, 50, 60]]


def test_new_io03_invalid_edge_drops_the_matching_count_in_every_series():
    result = read_length_boundary()._read_qscore({
        "bucket_ranges": [{"start": 0, "end": 1}, {"start": "bad", "end": 2}, {"start": 2, "end": 3}],
        "histogram_data": [{"bucket_values": [10, 20, 30]}, {"bucket_values": [40, 50, 60]}]})
    assert result.starts == [0.0, 2.0]
    assert result.ends == [1.0, 3.0]
    assert [series.values for series in result.series] == [[10, 30], [40, 60]]


@pytest.mark.parametrize("bad", [-10, -.2, "-0.2"])
def test_new_io04_negative_counts_are_absent_not_numeric_weights(bad):
    result = read_length_boundary()._read_buckets({
        "bucket_ranges": [{"start": 0, "end": 100}, {"start": 900, "end": 1100}],
        "histogram_data": [{"bucket_values": [bad, 20]}]})
    assert result.starts == [900]
    assert result.ends == [1100]
    assert result.values == [20]


@pytest.mark.parametrize("failure", ["write", "flush", "publish"])
def test_core01_atomic_stage_errors_preserve_original(tmp_path, monkeypatch, failure):
    output = tmp_path / "existing.html"
    output.write_text("ORIGINAL")
    if failure == "publish":
        monkeypatch.setattr(atomic_write, "_replace_with_retry", lambda *_: (_ for _ in ()).throw(OSError("publish")))
    elif failure == "flush":
        monkeypatch.setattr(atomic_write.os, "fsync", lambda *_: (_ for _ in ()).throw(OSError("flush")))
    with pytest.raises(OSError, match=failure):
        with atomic_write.atomic_output_path(output) as staged:
            staged.write_text("PARTIAL")
            if failure == "write":
                raise OSError("write")
    assert output.read_text(encoding="utf-8") == "ORIGINAL"
    assert list(tmp_path.iterdir()) == [output]


def test_core01_fallback_new_file_and_selected_html_write_succeed(tmp_path, monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda _: None)
    monkeypatch.setattr(importlib.util, "find_spec", lambda _: None)
    pdf = tmp_path / "report.pdf"
    result = pdf_boundary().export_pdf(object(), pdf)
    assert result["error"] is None
    html = pdf.with_suffix(".html")
    assert html.read_text(encoding="utf-8") == "<p>generic report</p>"
    html.write_text("OLD")
    result = pdf_boundary().export_pdf(object(), html)
    assert result["error"] is None
    assert html.read_text(encoding="utf-8") == "<p>generic report</p>"


def test_core01_pdf_module_converter_failure_keeps_previous(tmp_path, monkeypatch):
    output = tmp_path / "report.pdf"
    output.write_bytes(b"ORIGINAL")
    monkeypatch.setattr(shutil, "which", lambda _: None)
    monkeypatch.setattr(importlib.util, "find_spec", lambda _: object())
    def write_pdf(path):
        Path(path).write_bytes(b"PARTIAL")
        raise RuntimeError("module conversion failure")
    monkeypatch.setitem(sys.modules, "weasyprint", SimpleNamespace(HTML=lambda **_: SimpleNamespace(write_pdf=write_pdf)))
    result = pdf_boundary().export_pdf(object(), output)
    assert "module conversion failure" in result["error"]
    assert output.read_bytes() == b"ORIGINAL"
    assert list(tmp_path.iterdir()) == [output]


def test_core02_metadata_preserves_other_literal_string_prefixes(tmp_path):
    from kuma_core.shared.xlsx_text import append_literal_row
    values = ["=1+1", "+1", "-1", "@SUM(1)", "#REF!", " ordinary ", 7, False]
    wb = openpyxl.Workbook()
    sheet = wb.create_sheet("Metadata")
    append_literal_row(sheet, values)
    output = tmp_path / "literal.xlsx"
    wb.save(output)
    wb.close()
    with contextlib.closing(openpyxl.load_workbook(output)) as saved:
        assert [c.value for c in saved["Metadata"][1]] == values
        assert all(c.data_type == "s" for c in saved["Metadata"][1][:6])


def test_core03_well_formed_metadata_and_missing_timestamp_are_preserved(tmp_path):
    payload = {"protocol_run_info": {"flow_cell": {"flow_cell_id": "generic", "product_code": "product", "channel_count": 5}},
               "acquisitions": [{"acquisition_run_info": {"bream_info": {"mux_scan_results": [
                   {"counts": {"single_pore": 7}, "mux_scan_timestamp": 1.5},
                   {"counts": {"single_pore": 0}}]}}}]}
    (tmp_path / "report_generic.json").write_text(json.dumps(payload))
    result = load_metadata_module("flow_cell").read_flow_cell_history(tmp_path)
    assert (result.flow_cell_id, result.product_code, result.channel_count) == ("generic", "product", 5)
    assert [(s.at_seconds, s.single_pore) for s in result.scans] == [(1, 7), (0, 0)]


def test_core04_zip_member_count_is_bounded_before_parser(tmp_path):
    env = zip_boundary()
    env.load_structure_file.__globals__["MAX_ZIP_MEMBERS"] = 2
    path = tmp_path / "generic.zip"
    with zipfile.ZipFile(path, "w") as archive:
        for i in range(3):
            archive.writestr(f"generic-{i}.cif", " ")
    with pytest.raises(env.StructureFileError, match="member limit"):
        env.load_structure_file(path)


def test_core04_actual_stream_read_obeys_remaining_budget():
    env = zip_boundary()
    limits = []
    class Stream(io.BytesIO):
        def read(self, size: int | None = -1) -> bytes:
            limits.append(size)
            return super().read(size)
    archive = SimpleNamespace(open=lambda _: Stream(b" " * 10))
    with pytest.raises(env.StructureFileError, match="expanded-size limit"):
        env._read_zip_member(archive, "generic.cif", [8])
    assert limits == [9]
    budget = [20]
    assert len(env._read_zip_member(archive, "generic.cif", budget)) == 10
    assert budget == [10]


def test_core05_early_close_releases_source_generator():
    done = threading.Event()
    def source():
        try:
            yield from range(100)
        finally:
            done.set()
    env = boundary("kuma_core/mame/ingest/combinatorial_demux.py", {"_prefetch"})
    iterator = env._prefetch(source(), 1)
    assert next(iterator) == 0
    iterator.close()
    assert done.wait(1)


def test_core06_direct_ancestor_wins_and_ambiguous_siblings_are_absent(tmp_path):
    env = load_metadata_module("run_meta")
    setattr(env, "_MAX_LEVELS_UP", 1)
    # One level below tmp_path: a summary written at tmp_path itself makes this
    # test's directory a run folder beside every other test's tmp_path, and the
    # sibling search of any later discovery in the session then finds it.
    scope = tmp_path / "scope"
    input_dir = scope / "input"
    input_dir.mkdir(parents=True)
    for name in ["one", "two"]:
        folder = scope / name
        folder.mkdir()
        (folder / "final_summary_generic.txt").write_text(f"sample_id = {name}\n")
    assert env.discover_run_meta(input_dir) is None
    (scope / "final_summary_generic.txt").write_text("sample_id = parent\n")
    assert env.discover_run_meta(input_dir).sample_id == "parent"


def test_p03_preview_real_workbook_and_empty_sheet(tmp_path):
    wb = openpyxl.Workbook()
    sheet = wb.create_sheet("Data")
    sheet.append(["item", "value"])
    for i in range(10):
        sheet.append([f"item-{i}", i])
    path = tmp_path / "generic.xlsx"
    wb.save(path)
    wb.close()
    env = boundary("python-core/sidecar_kuro/handlers/misc.py", {"_preview_xlsx"})
    preview = env._preview_xlsx(str(path), "Data", 2)
    assert preview["headers"] == ["item", "value"]
    assert preview["rows"] == [["item-0", "0"], ["item-1", "1"]]
    empty = env._preview_xlsx(str(path), "Sheet", 2)
    assert empty["headers"] == [] and empty["rows"] == []
    with pytest.raises(ValueError, match="not found"):
        env._preview_xlsx(str(path), "missing", 2)


@pytest.mark.parametrize("raw", [True, False])
def test_p04_ledger_read_and_write_keep_source_identity(raw):
    tree = ast.parse((ROOT / "python-core/sidecar_mame/handlers/analyze.py").read_text(encoding="utf-8"))
    env = {"original_run_dir": Path("source"), "input_dir": Path("intermediate"), "is_raw": raw}
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
             and isinstance(n.func, ast.Name) and n.func.id in {"find_previous_use", "record_use"}]
    assert len(calls) == 2
    for call in calls:
        identity = eval(compile(ast.Expression(call.args[2]), "metadata_identity", "eval"), env)
        assert identity == ("source" if raw else "intermediate")


def test_new_io02_drop_reason_union_stays_in_sync():
    env = long_csv_boundary()
    from typing import get_args
    ts = (ROOT / "src/types/mame/activity.ts").read_text(encoding="utf-8")
    union = ts.split("export type DropReason =", 1)[1].split("\n\n", 1)[0]
    assert set(re.findall(r'"([^"]+)"', union)) == set(get_args(env.DropReason))


@pytest.mark.parametrize("values", [[1], [1, "bad"], [1, -3]])
def test_new_io03_invalid_only_qscore_series_is_absent(values):
    assert read_length_boundary()._read_qscore({
        "bucket_ranges": [{"start": 0, "end": 1}, {"start": 1, "end": 2}],
        "histogram_data": [{"bucket_values": values}]}) is None


def test_core01_html_report_handler_preserves_existing_on_publish_error(tmp_path, monkeypatch):
    from typing import cast
    output = tmp_path / "report.html"
    output.write_text("ORIGINAL")
    monkeypatch.setitem(sys.modules, "kuma_core.mame.report.builder", SimpleNamespace(
        build_run_report_data=lambda **_: SimpleNamespace()))
    monkeypatch.setitem(sys.modules, "kuma_core.mame.report.html_renderer", SimpleNamespace(
        render_html=lambda _: "<p>generic report</p>"))
    monkeypatch.setitem(sys.modules, "kuma_core.mame.report.pdf_export", SimpleNamespace(export_pdf=lambda *_: None))
    state = SimpleNamespace(last_verdicts=[object()], last_replicates=[], last_run_meta=None,
                            last_designed_mutant_ids=None)
    env = boundary("python-core/sidecar_mame/handlers/report.py", {"handle_export_run_report", "_ALLOWED_REPORT_EXTENSIONS"},
                   get_state=lambda: state, _validate_output_path=lambda path, **_: Path(path),
                   NgsRunMeta=object, cast=cast)
    monkeypatch.setattr(atomic_write, "_replace_with_retry", lambda *_: (_ for _ in ()).throw(OSError("publish failed")))
    with pytest.raises(RuntimeError, match="Failed to write HTML report"):
        env.handle_export_run_report({"output": str(output), "format": "html"})
    assert output.read_text(encoding="utf-8") == "ORIGINAL"
    assert list(tmp_path.iterdir()) == [output]


@pytest.mark.parametrize("handler", ["handle_save_json", "handle_save_workspace"])
def test_core01_json_success_and_nonfinite_rejection(tmp_path, handler):
    output = tmp_path / "state.json"
    env = boundary("python-core/sidecar_kuro/handlers/export.py", {handler},
                   SaveJsonParams=SimpleNamespace, SaveWorkspaceParams=SimpleNamespace,
                   FileExportResultModel=lambda **kw: SimpleNamespace(to_rpc_dict=lambda: kw),
                   _validate_output_path=lambda path, **_: Path(path))
    result = getattr(env, handler)({"filepath": str(output), "data": {"text": "일반", "value": 0}})
    assert result["filepath"] == str(output)
    assert json.loads(output.read_text(encoding="utf-8")) == {"text": "일반", "value": 0}
    before = output.read_bytes()
    with pytest.raises(ValueError):
        getattr(env, handler)({"filepath": str(output), "data": {"value": float("nan")}})
    assert output.read_bytes() == before


def test_core04_summary_member_stream_uses_same_total_budget():
    env = zip_boundary()
    class Archive:
        def namelist(self):
            return ["generic_summary_confidences_0.json"]
        def open(self, name):
            return io.BytesIO(b'{"ranking_score": 1}')
    with pytest.raises(env.StructureFileError, match="expanded-size limit"):
        env._summary_ranking_score(Archive(), "generic_model_0.cif", [4])


def test_core06_input_with_symlink_parent_is_not_followed(tmp_path):
    outside = tmp_path / "outside"
    outside.mkdir()
    deep = outside / "deep"
    deep.mkdir()
    (deep / "final_summary_generic.txt").write_text("sample_id = outside\n")
    link = tmp_path / "scope" / "link"
    link.parent.mkdir()
    link.symlink_to(outside, target_is_directory=True)
    assert load_metadata_module("run_meta").discover_run_meta(link / "deep") is None


def test_core07_derived_overflow_is_absent_and_finite_rows_recover(tmp_path):
    env = throughput_boundary()
    (tmp_path / "throughput_generic.csv").write_text(
        f"{env._TIME_COLS[0]},{env._CUMULATIVE_READS_COLS[0]}\n"
        "0,-1e308\n1,1e308\n2,1e308\n")
    result = env._parse_throughput(tmp_path)
    assert result == [{"time_h": .0333, "reads_per_sec": 0.0}]
    json.dumps(result, allow_nan=False)


def test_core01_encoding_failure_cleans_stage_and_keeps_original(tmp_path):
    output = tmp_path / "generic.json"
    output.write_text("ORIGINAL")
    with pytest.raises(UnicodeEncodeError):
        atomic_write.atomic_write_text(output, "bad surrogate \ud800")
    assert output.read_text(encoding="utf-8") == "ORIGINAL"
    assert list(tmp_path.iterdir()) == [output]


def test_core01_stage_reservation_failure_never_removes_existing_file(tmp_path, monkeypatch):
    stage = tmp_path / "other-file.pdf"
    stage.write_text("OTHER FILE")
    monkeypatch.setattr(atomic_write, "_temp_path_for", lambda _: stage)
    with pytest.raises(FileExistsError):
        with atomic_write.atomic_output_path(tmp_path / "report.pdf"):
            pytest.fail("reservation should fail")
    assert stage.read_text(encoding="utf-8") == "OTHER FILE"


@pytest.mark.skipif(atomic_write.os.name != "posix", reason="POSIX permission mode contract")
@pytest.mark.parametrize("mode", [0o600, 0o640])
@pytest.mark.parametrize("kind", ["text", "external"])
def test_core01_atomic_save_preserves_existing_private_permissions(tmp_path, mode, kind):
    output = tmp_path / "private.json"
    output.write_text("ORIGINAL")
    output.chmod(mode)
    old_umask = atomic_write.os.umask(0o022)
    try:
        if kind == "text":
            atomic_write.atomic_write_text(output, "NEW")
        else:
            with atomic_write.atomic_output_path(output) as stage:
                stage.write_text("NEW")
    finally:
        atomic_write.os.umask(old_umask)
    assert output.stat().st_mode & 0o777 == mode
    assert output.read_text(encoding="utf-8") == "NEW"


@pytest.mark.skipif(atomic_write.os.name != "posix", reason="POSIX permission mode contract")
def test_core01_external_stage_is_private_until_publish(tmp_path):
    output = tmp_path / "report.pdf"
    output.write_bytes(b"ORIGINAL")
    output.chmod(0o600)
    old_umask = atomic_write.os.umask(0o022)
    try:
        with atomic_write.atomic_output_path(output) as stage:
            assert stage.stat().st_mode & 0o777 == 0o600
            stage.write_bytes(b"generic")
    finally:
        atomic_write.os.umask(old_umask)
    assert output.stat().st_mode & 0o777 == 0o600


@pytest.mark.skipif(atomic_write.os.name != "posix", reason="POSIX permission mode contract")
@pytest.mark.parametrize("umask", [0o022, 0o002, 0o077])
@pytest.mark.parametrize("kind", ["text", "external"])
def test_core01_new_file_takes_the_mode_the_umask_gives(tmp_path, umask, kind):
    """A file this writer creates is as shareable as any other file the user
    creates, so a colleague on a shared folder can open it."""
    output = tmp_path / "new.json"
    old_umask = atomic_write.os.umask(umask)
    try:
        if kind == "text":
            atomic_write.atomic_write_text(output, "NEW")
        else:
            with atomic_write.atomic_output_path(output) as stage:
                stage.write_text("NEW")
    finally:
        atomic_write.os.umask(old_umask)
    assert output.stat().st_mode & 0o777 == 0o666 & ~umask
    assert output.read_text(encoding="utf-8") == "NEW"


def test_core01_text_stage_reservation_failure_never_removes_existing_file(tmp_path, monkeypatch):
    stage = tmp_path / "other-file.json"
    stage.write_text("OTHER FILE")
    monkeypatch.setattr(atomic_write, "_temp_path_for", lambda _: stage)
    with pytest.raises(FileExistsError):
        atomic_write.atomic_write_text(tmp_path / "report.json", "NEW")
    assert stage.read_text(encoding="utf-8") == "OTHER FILE"


@pytest.mark.skipif(atomic_write.os.name != "posix", reason="POSIX permission mode contract")
def test_core01_text_stage_is_private_before_publication(tmp_path, monkeypatch):
    output = tmp_path / "private.json"
    output.write_text("ORIGINAL")
    output.chmod(0o600)
    def inspect_stage(stage, target):
        assert stage.stat().st_mode & 0o777 == 0o600
        raise OSError("injected publish failure")
    monkeypatch.setattr(atomic_write, "_replace_with_retry", inspect_stage)
    old_umask = atomic_write.os.umask(0o022)
    try:
        with pytest.raises(OSError, match="injected"):
            atomic_write.atomic_write_text(output, "PRIVATE")
    finally:
        atomic_write.os.umask(old_umask)
    assert output.read_text(encoding="utf-8") == "ORIGINAL"
    assert list(tmp_path.iterdir()) == [output]


# ---------------------------------------------------------------------------
# core01: every user-named export publishes through a staged rename
# ---------------------------------------------------------------------------

def _mame_export_writers():
    """The MAME writers that publish to a path the operator named.

    Each entry is (id, callable taking the destination, previous bytes marker).
    Empty record lists keep these inert: no sequence design, analysis or
    instrument behaviour is exercised, only the file boundary.
    """
    from kuma_core.mame.activity.export_evolvepro import export_evolvepro_xlsx
    from kuma_core.mame.barcode_worklist import (
        BarcodeWorklist,
        write_barcode_worklist_csv,
    )
    from kuma_core.mame.export import (
        export_mame_janus_csv,
        export_mame_janus_xlsx,
        write_excel,
    )
    from kuma_core.mame.ingest import barcode_package
    from kuma_core.mame.io.variant_template import write_variant_template

    seeds = {
        f"{side}_{i}": "ACGTACGTAC"
        for side in ("fwd", "rev")
        for i in range(1, barcode_package._N_FWD + 1)
    }
    empty_worklist = BarcodeWorklist(
        rows=[], reverse_indices=[], forward_indices=[], missing_seeds=[]
    )
    return [
        ("janus_csv", lambda path: export_mame_janus_csv([], path), ".csv"),
        ("janus_xlsx", lambda path: export_mame_janus_xlsx([], path), ".xlsx"),
        ("result_workbook", lambda path: write_excel([], [], path), ".xlsx"),
        ("evolvepro_xlsx", lambda path: export_evolvepro_xlsx([], path), ".xlsx"),
        (
            "barcode_worklist",
            lambda path: write_barcode_worklist_csv(empty_worklist, path),
            ".csv",
        ),
        ("variant_template", lambda path: write_variant_template(path), ".xlsx"),
        (
            "package_barcodes_xlsx",
            lambda path: barcode_package._write_barcodes_xlsx(
                path=path,
                seeds=seeds,
                fwd_flanking="acgt",
                rev_flanking="acgt",
                gene_prefix="gene",
            ),
            ".xlsx",
        ),
        (
            "package_amplicon_fasta",
            lambda path: barcode_package._write_amplicon_fasta(
                path=path,
                cds_seq="ATG" * 40,
                gene_start=0,
                gene_end=120,
                gene_name="gene",
            ),
            ".fa",
        ),
        (
            "package_context_json",
            lambda path: barcode_package._write_mame_context_json(
                path=path,
                project_root=path.parent,
                barcodes_xlsx=path.parent / "barcodes_sequence.xlsx",
                amplicon_fa=path.parent / "gene_amplicon.fa",
            ),
            ".json",
        ),
    ]


@pytest.mark.parametrize(
    "name,writer,suffix",
    _mame_export_writers(),
    ids=[entry[0] for entry in _mame_export_writers()],
)
def test_core01_mame_export_publish_failure_keeps_previous_file(
    tmp_path, monkeypatch, name, writer, suffix
):
    """A Janus worklist or workbook that fails to publish leaves the old one.

    The lab reads these files as instructions. A truncated replacement that
    still carries the expected name is worse than no new file at all.
    """
    output = tmp_path / f"previous{suffix}"
    output.write_bytes(b"PREVIOUS RUN")
    monkeypatch.setattr(
        atomic_write,
        "_replace_with_retry",
        lambda *_: (_ for _ in ()).throw(OSError("publish")),
    )
    with pytest.raises(OSError, match="publish"):
        writer(output)
    assert output.read_bytes() == b"PREVIOUS RUN"
    assert list(tmp_path.iterdir()) == [output]


@pytest.mark.parametrize(
    "name,writer,suffix",
    _mame_export_writers(),
    ids=[entry[0] for entry in _mame_export_writers()],
)
def test_core01_mame_export_writes_a_sibling_stage_not_the_target(
    tmp_path, monkeypatch, name, writer, suffix
):
    """The bytes land on a sibling first; the target only ever appears whole."""
    output = tmp_path / f"previous{suffix}"
    output.write_bytes(b"PREVIOUS RUN")
    seen: list[tuple[Path, Path]] = []
    real_replace = atomic_write._replace_with_retry

    def record(tmp_target, final_target):
        seen.append((Path(tmp_target), Path(final_target)))
        assert Path(final_target) == output
        assert Path(tmp_target) != output
        assert Path(tmp_target).parent == output.parent
        assert output.read_bytes() == b"PREVIOUS RUN"
        real_replace(tmp_target, final_target)

    monkeypatch.setattr(atomic_write, "_replace_with_retry", record)
    writer(output)
    assert len(seen) == 1
    assert list(tmp_path.iterdir()) == [output]
    assert output.read_bytes() != b"PREVIOUS RUN"


def test_core01_text_stage_is_removed_when_the_content_cannot_be_written(tmp_path):
    """A non-OSError failure mid-write must not leave a .tmp in the export folder."""
    output = tmp_path / "generic.json"
    output.write_text("ORIGINAL")
    with pytest.raises(TypeError):
        atomic_write.atomic_write_text(output, 123)  # type: ignore[arg-type]
    assert output.read_text(encoding="utf-8") == "ORIGINAL"
    assert list(tmp_path.iterdir()) == [output]


def test_core01_text_stage_is_removed_when_the_write_is_interrupted(
    tmp_path, monkeypatch
):
    """Ctrl-C during the flush leaves the original file and no staged leftover."""
    output = tmp_path / "generic.json"
    output.write_text("ORIGINAL")
    monkeypatch.setattr(
        atomic_write.os,
        "fsync",
        lambda *_: (_ for _ in ()).throw(KeyboardInterrupt()),
    )
    with pytest.raises(KeyboardInterrupt):
        atomic_write.atomic_write_text(output, "NEW")
    assert output.read_text(encoding="utf-8") == "ORIGINAL"
    assert list(tmp_path.iterdir()) == [output]


# ---------------------------------------------------------------------------
# core01: the registry that keeps a new export handler from writing directly
# ---------------------------------------------------------------------------

# Writers a handler is allowed to delegate publication to, and the module that
# must carry the staging primitive for that delegation to mean anything.
# Writers a handler may delegate publication to, with the token whose presence
# in that module is the proof it stages. Most publish through
# ``atomic_output_path``; the Step 3 bundle publishes a whole set under one
# rollback, so its own primitive is the proof.
_ATOMIC_WRITERS = {
    "write_excel": ("kuma_core/mame/export/excel_writer.py", "atomic_output_path"),
    "export_mame_janus_csv": ("kuma_core/mame/export/janus_mapping.py", "atomic_output_path"),
    "export_mame_janus_xlsx": ("kuma_core/mame/export/janus_mapping.py", "atomic_output_path"),
    "export_evolvepro_xlsx": ("kuma_core/mame/activity/export_evolvepro.py", "atomic_output_path"),
    "export_pdf": ("kuma_core/mame/report/pdf_export.py", "atomic_output_path"),
    "write_barcode_worklist_csv": ("kuma_core/mame/barcode_worklist.py", "atomic_output_path"),
    "write_variant_template": ("kuma_core/mame/io/variant_template.py", "atomic_output_path"),
    "generate_mame_package": ("kuma_core/mame/ingest/barcode_package.py", "atomic_output_path"),
    "export_macrogen_xls": ("kuma_core/kuro/plate_mapper.py", "atomic_output_path"),
    "build_evolvepro_input": (
        "kuma_core/mame/activity/build_evolvepro_input.py",
        "_publish_artifact_bundle",
    ),
}

# Three anchors, not one. ``_validate_output_path`` alone missed
# ``handle_export_barcode_worklist`` and ``handle_export_variant_template``,
# which take the operator's path straight out of the params. A name anchor
# narrowed to export/save/generate missed nothing today and would miss a
# ``handle_write_*`` or ``handle_dump_*`` tomorrow, so it is every handler. The
# third anchor is the write itself: any function in these modules that opens a
# file for writing has to be registered, whatever it is called, which is what
# keeps a "no-write" declaration from being a free pass.
_HANDLER_NAME_ANCHOR = re.compile(r"^handle_")

_STAGING_NAMES = {"atomic_output_path", "atomic_write_text"}

# Calls that put bytes on disk without going through the staging primitives.
_WRITE_METHODS = {"write_text", "write_bytes", "save", "to_csv", "to_excel"}
_CSV_WRITER_FACTORIES = {"writer", "DictWriter"}
_WRITE_MODES = "wxa+"


def _write_mode_argument(call, name, is_attr):
    """Return (mode node, present) for an ``open``-family call.

    ``Path.open(mode, ...)`` puts the mode first while the builtin ``open`` and
    ``os.fdopen`` put the path or descriptor there, so the index depends on the
    shape of the call. Reading only ``args[1]`` let ``fasta_path.open("w")``
    through, which is how ``_write_reference_fasta`` looked inert.
    """
    index = 0 if (is_attr and name == "open") else 1
    if len(call.args) > index:
        return call.args[index], True
    for keyword in call.keywords:
        if keyword.arg == "mode":
            return keyword.value, True
    return None, False


def _direct_writes(node):
    """Return the direct-write calls inside *node*, as ``name@line`` strings.

    Modes are judged by exclusion rather than by a list of the ones known to
    write: anything containing ``w``, ``x``, ``a`` or ``+`` counts, and a mode
    that is not a literal string counts as well, because a variable mode cannot
    be proved harmless from here. The ``+`` cases (``r+``, ``rb+``, ``r+b``)
    write to an existing file and were invisible before. An ``open`` with no
    mode at all is the default read and stays quiet, which is what keeps the
    many ``open(path)`` reads in these modules from registering.

    The method list is NOT inverted the same way. A mode string is a closed
    grammar, so excluding is cheap and exact there, while treating every
    unknown method call as a possible write would put most of the repository
    into the registry for no gain in what the gate actually protects.

    ``.save``/``.write_text``/``.write_bytes``/``.to_csv``/``.to_excel`` and the
    two ``csv`` writer factories count as attribute calls, which is how an
    ``openpyxl`` workbook and a ``Path`` write are both seen without naming the
    receiver.
    """
    hits = []
    for call in ast.walk(node):
        if not isinstance(call, ast.Call):
            continue
        func = call.func
        is_attr = isinstance(func, ast.Attribute)
        name = func.attr if is_attr else (func.id if isinstance(func, ast.Name) else None)
        if name in ("open", "fdopen"):
            mode, present = _write_mode_argument(call, name, is_attr)
            if not present:
                continue
            if isinstance(mode, ast.Constant) and isinstance(mode.value, str):
                if any(ch in mode.value for ch in _WRITE_MODES):
                    hits.append(f"open({mode.value})@{call.lineno}")
            else:
                hits.append(f"open(<non-literal mode>)@{call.lineno}")
        elif is_attr and name in _WRITE_METHODS:
            hits.append(f".{name}()@{call.lineno}")
        elif is_attr and name in _CSV_WRITER_FACTORIES:
            hits.append(f"csv.{name}()@{call.lineno}")
    return hits


# An identifier that names an output destination. Narrower than "any path":
# ``filepath`` is the KURO name for a file being read as well as one being
# written, so matching it would flag every loader. ``output`` and the ``out_``
# forms are only ever destinations in this repository.
_OUTPUT_NAME = re.compile(r"output|out_path|out_dir|outdir|dest_path|dest_dir", re.I)

# Calls that convert or report rather than forward work. ``Path(p.output_dir)``
# and ``ValueError(f"output_dir must be ...")`` are not somebody else writing.
_NOT_A_FORWARD = {
    "Path", "str", "int", "float", "len", "bool", "isinstance", "getattr",
    "print", "format", "repr", "sorted", "ValueError", "RuntimeError",
    "FileNotFoundError", "TypeError", "KeyError", "NotADirectoryError",
    "debug", "info", "warning", "error", "exception",
    # Validating a path reads it. Handing the operator's path to the validator
    # is what every export handler does before it decides anything, so counting
    # it as delegation would make the rule fire on the check itself.
    "_validate_output_path", "_validate_filepath",
}


# How a handler usually gets hold of the operator's path.
_PATH_VALIDATORS = {"_validate_output_path"}
_MAPPING_READERS = {"get", "pop"}


def _output_name_in(expr, aliases=frozenset()):
    """Name an output destination mentioned inside *expr*, or None.

    Bare string constants are ignored so that a message reading "output_dir
    must be absolute" is not mistaken for a destination. A string that is a
    subscript key is kept, because ``params["output"]`` is exactly how a
    handler takes the operator's path when it skips the model.

    *aliases* are local names :func:`_output_path_aliases` proved hold a
    destination even though they are not named like one.
    """
    for node in ast.walk(expr):
        if isinstance(node, ast.Name) and (
            _OUTPUT_NAME.search(node.id) or node.id in aliases
        ):
            return node.id
        if isinstance(node, ast.Attribute) and _OUTPUT_NAME.search(node.attr):
            return node.attr
        if isinstance(node, ast.Subscript):
            key = node.slice
            if (
                isinstance(key, ast.Constant)
                and isinstance(key.value, str)
                and _OUTPUT_NAME.search(key.value)
            ):
                return f"[{key.value!r}]"
    return None


def _output_path_aliases(node):
    """Local names that hold an output destination under another name.

    The name rule alone reads the call site, and KURO binds the operator's path
    to ``resolved``, which looks like nothing. ``resolved =
    _validate_output_path(p.filepath)`` followed by ``export_idt_csv(results,
    resolved)`` therefore showed no forward at all, and a "no-write"
    declaration on such a handler passed. Three bindings count:

    - the result of a path validator, which is what the value is for
    - an expression the name rule already matches, so ``dest = p.output_xlsx``
      carries over
    - ``params.get("output_path")`` and friends, the mapping form of the same

    One level, no transitive propagation through further assignments. The point
    is to see the operator's path leave the function, and one binding is what
    stands between the validator and the writer in every case here.

    Position does not matter: ``ast.walk`` collects every binding in the
    function before any call is examined, so an alias assigned below the call
    that uses it still counts. The error that leaves is over-detection rather
    than under-detection, and a spurious forward costs a registry label rather
    than a missed write.
    """
    aliases = set()
    for statement in ast.walk(node):
        if isinstance(statement, ast.Assign):
            targets = statement.targets
        elif isinstance(statement, ast.AnnAssign):
            targets = [statement.target]
        else:
            continue
        if len(targets) != 1 or not isinstance(targets[0], ast.Name):
            continue
        value = statement.value
        if value is None:
            continue
        if _binds_an_output_path(value):
            aliases.add(targets[0].id)
    # Frozen on the way out, not at each use: the only consumer tests
    # membership and never mutates, which the parameter type states, and
    # freezing here converts once per handler rather than once per argument.
    return frozenset(aliases)


def _binds_an_output_path(value):
    for node in ast.walk(value):
        if not isinstance(node, ast.Call):
            continue
        func = node.func
        name = func.attr if isinstance(func, ast.Attribute) else getattr(func, "id", None)
        if name in _PATH_VALIDATORS:
            return True
        if (
            name in _MAPPING_READERS
            and node.args
            and isinstance(node.args[0], ast.Constant)
            and isinstance(node.args[0].value, str)
            and _OUTPUT_NAME.search(node.args[0].value)
        ):
            return True
    return _output_name_in(value) is not None


def _forwarded_output_paths(node):
    """Return the calls in *node* that hand an output destination onward.

    The gate used to read the handler body alone, so a handler that took the
    operator's path and passed it to a writer one call away declared itself
    inert and passed. Three did: the Step 3 builder, the combinatorial demux
    and the KURO export-everything fan-out. Following the whole call graph is
    not needed to see it. What the operator's path is named, and the fact that
    it leaves the function as an argument, is enough and stays cheap.
    """
    aliases = _output_path_aliases(node)
    hits = []
    for call in ast.walk(node):
        if not isinstance(call, ast.Call):
            continue
        func = call.func
        callee = (
            func.attr if isinstance(func, ast.Attribute)
            else func.id if isinstance(func, ast.Name)
            else "<expr>"
        )
        if callee in _NOT_A_FORWARD:
            continue
        for argument in list(call.args) + [kw.value for kw in call.keywords]:
            found = _output_name_in(argument, aliases)
            if found:
                hits.append(f"{callee}(...{found}...)@{call.lineno}")
                break
        else:
            for keyword in call.keywords:
                if keyword.arg and _OUTPUT_NAME.search(keyword.arg):
                    hits.append(f"{callee}({keyword.arg}=)@{call.lineno}")
                    break
    return hits


# Every function in a sidecar handler module that one of the three anchors
# reaches.
# "here"          the function body stages the write itself
# "callee"        it hands the resolved path to a writer from _ATOMIC_WRITERS
# "transitive"    the writer is reached through the named intermediate module
# "validate-only" the function checks a path and its own body writes nothing
# "no-write"      its own body performs no direct write and forwards no
#                 destination (asserted, not taken on trust). It does NOT
#                 say nothing is written anywhere downstream: a write
#                 reached through a helper the anchors do not see is
#                 invisible here, which is why three handlers wore this
#                 label while installing a codon table or a structure cache
# "app-state"     application state under the user profile or a cache is
#                 written, by the body or by a helper it calls
# "exempt-write"  it writes directly to a path the operator never named
# "direct"        known debt: an operator path written without a staged rename
#
# A new function any anchor reaches is absent from this table and fails the
# test below, which is the point: it has to state how it publishes. The two
# kinds that admit a direct write, "direct" and "exempt-write", are each also
# frozen in a set further down, so declaring one is a two-place edit a reviewer
# sees rather than a word changed on one line.
_OUTPUT_PATH_HANDLERS = {
    ("sidecar_kuro/handlers/codon.py", "handle_import_codon_table"): ("app-state", None),
    ("sidecar_kuro/handlers/codon.py", "handle_export_codon_table"): ("here", None),
    ("sidecar_kuro/handlers/codon.py", "handle_compute_codon_table"): ("app-state", None),
    ("sidecar_kuro/handlers/design.py", "handle_design_sdm_primers"): ("exempt-write", None),
    ("sidecar_kuro/handlers/design.py", "handle_retry_failed"): ("no-write", None),
    ("sidecar_kuro/handlers/design.py", "handle_swap_primer"): ("no-write", None),
    ("sidecar_kuro/handlers/design.py", "handle_commit_design_result"): ("no-write", None),
    ("sidecar_kuro/handlers/design.py", "handle_evaluate_primer"): ("no-write", None),
    ("sidecar_kuro/handlers/design.py", "handle_get_alternatives"): ("no-write", None),
    ("sidecar_kuro/handlers/export.py", "handle_get_plate_map"): ("no-write", None),
    ("sidecar_kuro/handlers/export.py", "handle_export_excel"): ("here", None),
    ("sidecar_kuro/handlers/export.py", "handle_export_order"): ("transitive",
        ("kuma_core/kuro/plate_mapper.py", "atomic_output_path")),
    ("sidecar_kuro/handlers/export.py", "handle_export_mapping"): ("transitive",
        ("kuma_core/kuro/plate_mapper.py", "atomic_output_path")),
    ("sidecar_kuro/handlers/export.py", "handle_save_workspace"): ("here", None),
    ("sidecar_kuro/handlers/export.py", "handle_save_json"): ("here", None),
    ("sidecar_kuro/handlers/export.py", "handle_load_workspace"): ("no-write", None),
    ("sidecar_kuro/handlers/export.py", "handle_export_benchmark_csv"): ("here", None),
    ("sidecar_kuro/handlers/export.py", "handle_export_macrogen"): ("callee",
        "export_macrogen_xls"),
    ("sidecar_kuro/handlers/export.py", "handle_export_echo_mapping_dry_run"): ("no-write", None),
    ("sidecar_kuro/handlers/export.py", "handle_export_janus_mapping_dry_run"): ("no-write", None),
    # Six of the eight bundle artefacts are written by plate_mapper exporters,
    # which is what this points at. The other two go through
    # ``_export_primers_fasta`` in this same module and ``write_run_manifest``
    # (``kuma_core/shared/run_manifest.py``); both call ``atomic_write_text``.
    # The kind names one intermediate by design, so the rest is this comment.
    # ``_export_primers_fasta`` is deliberately not registered: it staged its
    # way out of every anchor in ``_anchored``, and a write put back into it
    # returns it to the found set as "unregistered", which is the gate working.
    ("sidecar_kuro/handlers/export.py", "handle_export_all"): ("transitive",
        ("kuma_core/kuro/plate_mapper.py", "atomic_output_path")),
    ("sidecar_kuro/handlers/external.py", "handle_fetch_domains"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_search_uniprot"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_check_structures_available"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_fetch_structure"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_load_structure_file"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_fetch_interface_residues"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_fetch_pdb_text"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_fetch_active_site"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_compute_dispersion"): ("no-write", None),
    ("sidecar_kuro/handlers/external.py", "handle_predict_structure_esmfold"): ("app-state", None),
    ("sidecar_kuro/handlers/external.py", "handle_annotate_domains_by_sequence"): ("no-write", None),
    ("sidecar_kuro/handlers/misc.py", "handle_list_polymerases"): ("no-write", None),
    ("sidecar_kuro/handlers/misc.py", "handle_get_polymerase_details"): ("no-write", None),
    ("sidecar_kuro/handlers/misc.py", "handle_save_custom_polymerase"): ("app-state", None),
    ("sidecar_kuro/handlers/misc.py", "ensure_user_codon_dir"): ("app-state", None),
    ("sidecar_kuro/handlers/misc.py", "handle_list_organisms"): ("no-write", None),
    ("sidecar_kuro/handlers/misc.py", "handle_preview_evolvepro_source"): ("no-write", None),
    ("sidecar_kuro/handlers/misc.py", "handle_load_evolvepro_csv"): ("no-write", None),
    ("sidecar_kuro/handlers/misc.py", "handle_run_benchmark"): ("no-write", None),
    ("sidecar_kuro/handlers/sequence.py", "handle_load_fasta"): ("no-write", None),
    ("sidecar_kuro/handlers/sequence.py", "handle_parse_mutations_text"): ("no-write", None),
    ("sidecar_kuro/handlers/settings.py", "handle_load"): ("no-write", None),
    ("sidecar_kuro/handlers/settings.py", "handle_save"): ("app-state", None),  # writes via os.fdopen
    ("sidecar_mame/handlers/activity.py", "handle_activity_upload"): ("no-write", None),
    ("sidecar_mame/handlers/activity.py", "handle_activity_set_plate_meta"): ("no-write", None),
    ("sidecar_mame/handlers/activity.py", "handle_activity_merge"): ("no-write", None),
    ("sidecar_mame/handlers/activity.py", "handle_activity_export_evolvepro_xlsx"): ("callee",
        "export_evolvepro_xlsx"),
    ("sidecar_mame/handlers/activity.py", "handle_merge_for_evolvepro"): ("no-write", None),
    ("sidecar_mame/handlers/activity.py", "handle_build_evolvepro_input"): ("callee", "build_evolvepro_input"),
    ("sidecar_mame/handlers/activity.py", "handle_detect_measurement_source"): ("no-write", None),
    ("sidecar_mame/handlers/analyze.py", "_acceptance_findings"): ("validate-only", None),
    ("sidecar_mame/handlers/analyze.py", "_write_reference_fasta"): ("exempt-write", None),
    ("sidecar_mame/handlers/analyze.py", "handle_validate_inputs"): ("no-write", None),
    ("sidecar_mame/handlers/analyze.py", "handle_analyze"): ("transitive",
        ("kuma_core/mame/pipeline.py", "write_excel")),
    ("sidecar_mame/handlers/barcode_package.py", "handle_inspect_variant_source"): ("no-write", None),
    ("sidecar_mame/handlers/barcode_package.py", "handle_check_plate_order"): ("no-write", None),
    ("sidecar_mame/handlers/barcode_package.py", "handle_generate_mame_package"): ("callee",
        "generate_mame_package"),
    ("sidecar_mame/handlers/barcode_worklist.py", "handle_export_barcode_worklist"): ("callee",
        "write_barcode_worklist_csv"),
    ("sidecar_mame/handlers/build_well_layout.py", "handle_build_well_layout"): ("no-write", None),
    ("sidecar_mame/handlers/classify_round.py", "handle_classify_round"): ("no-write", None),
    ("sidecar_mame/handlers/combinatorial_demux.py", "handle_run_combinatorial_demux"): ("transitive", ("kuma_core/mame/ingest/combinatorial_demux.py", "atomic_write_text")),
    ("sidecar_mame/handlers/demux.py", "handle_demux_and_filter"): ("here", None),
    ("sidecar_mame/handlers/detect_native_barcodes.py", "handle_detect_native_barcodes"): ("no-write", None),
    ("sidecar_mame/handlers/export.py", "handle_export_excel"): ("callee", "write_excel"),
    ("sidecar_mame/handlers/export.py", "handle_get_plate_data"): ("no-write", None),
    ("sidecar_mame/handlers/export.py", "handle_apply_selection_overrides"): ("transitive", ("python-core/sidecar_mame/handlers/export.py", "write_excel")),
    ("sidecar_mame/handlers/export.py", "handle_export_janus_mapping"): ("callee",
        "export_mame_janus_csv"),
    ("sidecar_mame/handlers/export.py", "handle_export_janus_mapping_dry_run"): ("no-write", None),
    ("sidecar_mame/handlers/health.py", "handle_get_run_health"): ("no-write", None),
    ("sidecar_mame/handlers/ingest.py", "handle_parse_reference"): ("no-write", None),
    ("sidecar_mame/handlers/kuma_meta.py", "handle_read_kuma_meta"): ("no-write", None),
    ("sidecar_mame/handlers/load.py", "handle_load_analyze_result"): ("transitive",
        ("python-core/sidecar_mame/handlers/export.py", "write_excel")),
    ("sidecar_mame/handlers/report.py", "handle_export_run_report"): ("here", None),
    ("sidecar_mame/handlers/variant_template.py", "handle_export_variant_template"): ("callee",
        "write_variant_template"),
}

# Handlers that still write the operator's path directly. The set may shrink,
# never grow: a new "direct" entry fails this assertion.
#
# Empty today. It held the eight KURO export handlers until they were staged
# the way the MAME ones were, so the sidecars now agree that no operator-named
# output is written in place. Kept rather than deleted, because an empty frozen
# set is what makes the next "direct" registration a visible set edit instead
# of one word on one line.
_KNOWN_DIRECT_WRITERS: set[tuple[str, str]] = set()

# Functions that write directly to a path the operator never chose. Frozen for
# the same reason as the debt set above.
#   handle_design_sdm_primers  a NamedTemporaryFile holding the mutation list
#   _write_reference_fasta     a derived FASTA inside the run's own demux
#                              output directory or a TemporaryDirectory
#   ensure_user_codon_dir      the bundled seed tables copied into that same
#                              user directory on first use
#   handle_save                preferences.json under the user profile, staged
#                              with mkstemp and renamed
_KNOWN_NON_OPERATOR_WRITERS = {
    ("sidecar_kuro/handlers/design.py", "handle_design_sdm_primers"),
    ("sidecar_mame/handlers/analyze.py", "_write_reference_fasta"),
    ("sidecar_kuro/handlers/misc.py", "ensure_user_codon_dir"),
    ("sidecar_kuro/handlers/settings.py", "handle_save"),
}

# Handlers that pass an output destination to something that does not write it.
#
# What this kind guarantees: the body performs no direct write, it does forward
# a destination, and a human put the entry in this set.
#
# What it does NOT guarantee: that the receiving function leaves the path
# alone. Nothing here reads the callee. The label is only as good as the review
# that added the member, which is why it is a set edit and not a word.
#
# Empty today. ``handle_load_analyze_result`` held the one entry until the path
# it forwards turned out to reach ``_rewrite_result_workbook`` two RPCs later,
# so it is declared ``transitive`` instead. Kept because the shape recurs and
# the next one should land here rather than in ``no-write``.
_KNOWN_NON_WRITING_FORWARDERS: set[tuple[str, str]] = set()


class _HandlerFacts(NamedTuple):
    calls: set
    writes: list
    forwards: list


def _handler_functions(source, origin="<registry>"):
    """Map every function in *source* to the names it calls and the writes it does.

    Attribute calls are included alongside plain names, so ``mod.write_excel(...)``
    after an ``import`` of the module, or an aliased import, reads the same as a
    bare call. Reading only ``ast.Name`` let both forms through unseen.
    """
    found = {}
    for node in ast.walk(ast.parse(source, filename=origin)):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            # Two definitions of one name in one module would collapse into
            # the later one and take the earlier one's writes with them. No
            # module does that today; the check keeps it that way.
            assert node.name not in found, (origin, node.name)
            found[node.name] = _HandlerFacts(
                calls={
                    (call.func.attr if isinstance(call.func, ast.Attribute) else call.func.id)
                    for call in ast.walk(node)
                    if isinstance(call, ast.Call)
                    and isinstance(call.func, (ast.Name, ast.Attribute))
                },
                writes=_direct_writes(node),
                forwards=_forwarded_output_paths(node),
            )
    return found


def _handler_modules():
    """Every handler module, including any future subpackage.

    The handler directories are found first and recursed into second, so a
    subpackage added under one of them is covered while a build or vendor tree
    elsewhere under ``python-core`` cannot match. A flat glob would silently
    stop covering the day a sidecar grows a subpackage. ``__pycache__`` is
    skipped because its ``.py`` files do not exist.
    """
    root = ROOT / "python-core"
    return sorted(
        p.relative_to(root).as_posix()
        for handlers in root.glob("sidecar_*/handlers")
        for p in handlers.rglob("*.py")
        if "__pycache__" not in p.parts
    )


def _anchored(name, facts):
    """Three anchors. Forwarding is NOT one of them.

    Forwarding an output path constrains what a registered handler may claim,
    and making it an anchor as well would pull in every private helper that
    passes a destination along, which is noise the gate does not need.
    """
    return bool(
        _HANDLER_NAME_ANCHOR.match(name)
        or "_validate_output_path" in facts.calls
        or facts.writes
    )


def _handlers_resolving_user_paths(modules=None, read=None):
    modules = _handler_modules() if modules is None else modules
    read = read or (lambda m: (ROOT / "python-core" / m).read_text(encoding="utf-8"))
    return {
        (module, name): facts
        for module in modules
        for name, facts in _handler_functions(read(module), module).items()
        if _anchored(name, facts)
    }


def test_core01_every_user_output_handler_declares_how_it_publishes():
    """A handler that resolves an operator path must be registered above.

    Found by AST rather than by grep so that renaming a handler, or adding a
    new one in either sidecar, cannot slip past unlisted.
    """
    _check_registry(_handlers_resolving_user_paths())


def _check_registry(found, registry=None):
    registry = _OUTPUT_PATH_HANDLERS if registry is None else registry
    assert set(found) == set(registry), {
        "unregistered": sorted(set(found) - set(registry)),
        "stale": sorted(set(registry) - set(found)),
    }
    for key, facts in found.items():
        kind, detail = registry[key]
        _check_one_handler(key, facts, kind, detail)


def _check_one_handler(key, facts, kind, detail):
    """Verify one registered handler against what its source actually does."""
    if kind == "here":
        assert facts.calls & _STAGING_NAMES, key
    elif kind == "callee":
        assert detail in facts.calls, key
        _assert_writer_stages(detail, key)
    elif kind == "transitive":
        intermediate, token = detail
        assert token in (ROOT / intermediate).read_text(encoding="utf-8"), key
        if token in _ATOMIC_WRITERS:
            _assert_writer_stages(token, key)
        else:
            # No named writer to point at, so the intermediate module has
            # to carry a staging primitive itself.
            source = (ROOT / intermediate).read_text(encoding="utf-8")
            assert any(n in source for n in _STAGING_NAMES), (key, intermediate)
    elif kind in ("validate-only", "no-write"):
        # The declaration is checked, not believed. A body that opens a
        # file for writing cannot wear one of these labels, and neither can
        # one that hands an output destination to somebody else: that was
        # the hole the first version of this gate left open.
        assert not facts.writes, (key, kind, facts.writes)
        assert not facts.calls & _STAGING_NAMES, key
        assert not facts.forwards, (key, kind, facts.forwards)
    elif kind == "forwards-only":
        # Passes a destination onward to something that does not write it,
        # such as the in-memory analyze cache. Frozen, so the claim is a
        # reviewed edit rather than a label.
        assert facts.forwards, key
        assert not facts.writes, (key, facts.writes)
        assert key in _KNOWN_NON_WRITING_FORWARDERS, key
    elif kind in ("app-state", "exempt-write"):
        # These may write, to a path the operator never chose. Saying so
        # costs a second edit in the frozen set below, which a reviewer
        # sees, rather than one word changed on one line.
        assert not facts.writes or key in _KNOWN_NON_OPERATOR_WRITERS, (key, facts.writes)
    else:
        # Debt. Writing the operator's path straight is allowed only for
        # the handlers already named in the frozen set below.
        assert kind == "direct", key
        assert key in _KNOWN_DIRECT_WRITERS, key


def _assert_writer_stages(writer, key):
    module, token = _ATOMIC_WRITERS[writer]
    assert token in (ROOT / module).read_text(encoding="utf-8"), (key, module, token)


def test_core01_direct_writer_debt_may_shrink_but_never_grow():
    direct = {
        key for key, (kind, _) in _OUTPUT_PATH_HANDLERS.items() if kind == "direct"
    }
    assert direct <= _KNOWN_DIRECT_WRITERS, sorted(direct - _KNOWN_DIRECT_WRITERS)


def test_core01_delegated_writers_actually_stage_their_output():
    """The delegation targets are only trustworthy if they really stage."""
    for writer, (module, token) in _ATOMIC_WRITERS.items():
        source = (ROOT / module).read_text(encoding="utf-8")
        assert f"def {writer}(" in source, writer
        assert token in source, (module, token)


def test_core01_the_frozen_sets_hold_no_entry_the_registry_dropped():
    """A frozen exemption outliving its registry entry is a dead licence.

    ``_write_reference_fasta`` became one the moment the write detector was
    thought to have seen it and had not, and nothing said so.
    """
    registered = set(_OUTPUT_PATH_HANDLERS)
    for name, frozen in (
        ("_KNOWN_DIRECT_WRITERS", _KNOWN_DIRECT_WRITERS),
        ("_KNOWN_NON_OPERATOR_WRITERS", _KNOWN_NON_OPERATOR_WRITERS),
        ("_KNOWN_NON_WRITING_FORWARDERS", _KNOWN_NON_WRITING_FORWARDERS),
    ):
        assert frozen <= registered, (name, sorted(frozen - registered))


# Two handlers that would slip past a gate that trusted its declaration. The
# first writes in its own body. The second is the shape the first version of
# this gate missed entirely: the body is clean and the write is one call away.
_BYPASS_FIXTURES = {
    "direct": """
def handle_export_something(params):
    output = params["output"]
    with open(output, "w", encoding="utf-8") as fh:
        fh.write("partial")
    return {"output_path": output}
""",
    "delegating": """
def _really_write_it(output_path):
    with open(output_path, "w", encoding="utf-8") as fh:
        fh.write("partial")


def handle_export_something(params):
    output = params["output"]
    _really_write_it(output)
    return {"output_path": output}
""",
}


def _fixture_registry(kind):
    return {("sidecar_mame/handlers/fixture.py", "handle_export_something"): (kind, None)}


def _fixture_found(shape="direct"):
    return _handlers_resolving_user_paths(
        modules=["sidecar_mame/handlers/fixture.py"],
        read=lambda _: _BYPASS_FIXTURES[shape],
    )


@pytest.mark.parametrize("kind", ["no-write", "validate-only"])
def test_core01_a_quiet_declaration_cannot_hide_a_direct_write(kind):
    """Declaring a writing handler inert must fail, not pass.

    Before the write check this was the cheapest way past the registry: one
    word on one line, with ``open(path, "w")`` still in the body.
    """
    found = _fixture_found()
    assert list(found) == [("sidecar_mame/handlers/fixture.py", "handle_export_something")]
    with pytest.raises(AssertionError) as caught:
        _check_registry(found, _fixture_registry(kind))
    assert "open(w)@4" in str(caught.value)


@pytest.mark.parametrize("kind", ["no-write", "validate-only"])
def test_core01_a_quiet_declaration_cannot_hide_a_delegated_write(kind):
    """The body is clean and the operator's path still reaches ``open(..., "w")``.

    This is the shape that got past the first version of the gate three times
    over: the Step 3 builder, the combinatorial demux and the KURO
    export-everything fan-out all declared themselves inert while handing the
    destination to the function that writes it.
    """
    handler = ("sidecar_mame/handlers/fixture.py", "handle_export_something")
    facts = _fixture_found("delegating")[handler]
    # The body itself is clean. Only the forward gives it away.
    assert facts.writes == []
    with pytest.raises(AssertionError) as caught:
        _check_one_handler(handler, facts, kind, None)
    assert "_really_write_it(...output...)" in str(caught.value)


def test_core01_an_unregistered_writing_handler_fails_the_gate():
    with pytest.raises(AssertionError) as caught:
        _check_registry(_fixture_found(), {})
    assert "unregistered" in str(caught.value)


def test_core01_an_exempt_write_must_also_be_named_in_the_frozen_set():
    with pytest.raises(AssertionError):
        _check_registry(_fixture_found(), _fixture_registry("exempt-write"))


def test_core01_a_direct_declaration_must_also_be_named_in_the_debt_set():
    with pytest.raises(AssertionError):
        _check_registry(_fixture_found(), _fixture_registry("direct"))


def test_core01_the_write_detector_sees_every_shape_it_claims_to():
    """Verified against inputs whose answer is known, both ways.

    The reads must stay absent and the writes must all appear, including the
    two shapes that used to escape: ``Path.open`` with the mode as its first
    argument, and a mode that updates in place rather than truncating.
    """
    writes = _direct_writes(
        ast.parse(
            "def f(p, fd, fh, mode):\n"
            "    open(p)\n"
            "    open(p, 'r')\n"
            "    p.open()\n"
            "    p.open('rb')\n"
            "    p.read_text()\n"
            "    open(p, 'w')\n"
            "    os.fdopen(fd, 'w')\n"
            "    open(p, mode='ab')\n"
            "    p.open('w', encoding='utf-8')\n"
            "    p.open('r+b')\n"
            "    open(p, 'rb+')\n"
            "    open(p, mode)\n"
            "    wb.save(p)\n"
            "    p.write_text('x')\n"
            "    p.write_bytes(b'x')\n"
            "    df.to_csv(p)\n"
            "    df.to_excel(p)\n"
            "    csv.writer(fh)\n"
            "    csv.DictWriter(fh, [])\n"
        ).body[0]
    )
    assert [w.split("@")[0] for w in writes] == [
        "open(w)",
        "open(w)",
        "open(ab)",
        "open(w)",
        "open(r+b)",
        "open(rb+)",
        "open(<non-literal mode>)",
        ".save()",
        ".write_text()",
        ".write_bytes()",
        ".to_csv()",
        ".to_excel()",
        "csv.writer()",
        "csv.DictWriter()",
    ]


def test_core01_the_forward_detector_separates_delegation_from_reporting():
    """Known-answer check on the rule that closed the call-chain hole."""
    forwards = _forwarded_output_paths(
        ast.parse(
            "def f(params, p):\n"
            "    raise_if(ValueError('output_dir must be absolute'))\n"
            "    logger.info('wrote output=%s', 1)\n"
            "    Path(p.output_dir)\n"
            "    _validate_output_path(params['output'])\n"
            "    read_variant_source(p.expected_mutations_xlsx)\n"
            "    load_fasta(p.filepath)\n"
            "    build_it(p.output_xlsx)\n"
            "    run_it(output_dir=p.somewhere)\n"
            "    write_it(params['output_path'])\n"
            "    resolved = _validate_output_path(p.filepath)\n"
            "    export_idt_csv(results, resolved)\n"
            "    dest = p.output_xlsx\n"
            "    save_it(dest)\n"
            "    picked = params.get('output_dir')\n"
            "    dump_it(picked)\n"
        ).body[0]
    )
    assert [f.split("@")[0] for f in forwards] == [
        "build_it(...output_xlsx...)",
        "run_it(output_dir=)",
        "write_it(...['output_path']...)",
        # The three alias shapes. Without them a handler could bind the
        # operator's path to a name of its own and hand it to a writer with
        # nothing in the call site to show for it, which is how KURO spells
        # every one of its exports.
        "export_idt_csv(...resolved...)",
        "save_it(...dest...)",
        "dump_it(...picked...)",
    ]



# ---------------------------------------------------------------------------
# core01: the hand-rolled stage-and-rename writers clean up for any failure
# ---------------------------------------------------------------------------

def test_core01_bundle_rollback_survives_an_interrupt_between_publishes(
    tmp_path, monkeypatch
):
    """A Ctrl-C between two publishes leaves the whole bundle on its old version.

    Each file's publish is an atomic replace, so no single file is ever half
    written. What used to escape the rollback was the gap between two of them.
    The interrupt here lands after the first destination already holds the new
    bytes and before the second does, which is the state the rollback exists to
    undo, and the assertion below proves that state was reached.
    """
    # The module, not the same-named function the package re-exports.
    from kuma_core.mame.activity.build_evolvepro_input import (
        _publish_artifact_bundle,
    )

    first = tmp_path / "round1.xlsx"
    second = tmp_path / "gc_export.xlsx"
    manifest = tmp_path / "round1.xlsx.manifest.json"
    first.write_text("OLD FIRST")
    second.write_text("OLD SECOND")
    manifest.write_text("OLD MANIFEST")

    writers = [
        (first, lambda path: path.write_text("NEW FIRST")),
        (second, lambda path: path.write_text("NEW SECOND")),
    ]

    real_replace = Path.replace
    published: list[str] = []

    def interrupt_between_publishes(self, target):
        # Chosen by destination rather than by call number, so a change to how
        # many files the bundle moves aside first cannot silently move the
        # injection point. The first publish goes through and the second is
        # interrupted, which is the only moment the set is half new and half
        # old. Stopping before any publish would prove nothing: the rollback
        # would have had nothing to undo.
        if Path(target) == second and first.read_text(encoding="utf-8") == "NEW FIRST":
            raise KeyboardInterrupt()
        result = real_replace(self, target)
        published.append(Path(target).name)
        return result

    monkeypatch.setattr(Path, "replace", interrupt_between_publishes)
    with pytest.raises(KeyboardInterrupt):
        _publish_artifact_bundle(writers, manifest, {"schema": 1})

    # The first destination really did receive its new bytes before the
    # interrupt, so what follows is the rollback undoing a publish rather than
    # a bundle that never started.
    assert first.name in published

    assert first.read_text(encoding="utf-8") == "OLD FIRST"
    assert second.read_text(encoding="utf-8") == "OLD SECOND"
    assert manifest.read_text(encoding="utf-8") == "OLD MANIFEST"
    assert sorted(p.name for p in tmp_path.iterdir()) == sorted(
        [first.name, second.name, manifest.name]
    )


def test_core01_preferences_save_leaves_no_stage_when_interrupted(
    tmp_path, monkeypatch
):
    """The preferences writer is a second hand-rolled stage-and-rename.

    It cleaned up for OSError only, so an interrupt between the write and the
    rename left a ``preferences_*.tmp`` beside the real file.
    """
    from sidecar_kuro.handlers import settings as settings_handler

    prefs = tmp_path / "preferences.json"
    monkeypatch.setenv("KUMA_PREFERENCES_PATH", str(prefs))
    monkeypatch.setattr(
        settings_handler.os,
        "replace",
        lambda *_: (_ for _ in ()).throw(KeyboardInterrupt()),
    )
    with pytest.raises(KeyboardInterrupt):
        settings_handler.handle_save({"settings": {"theme": "dark"}})
    assert not prefs.exists()
    assert list(tmp_path.iterdir()) == []
